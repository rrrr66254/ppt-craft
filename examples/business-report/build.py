"""Example deck: a small coffee roastery's quarterly review (report-grid preset, Korean).

The company and every number are fictional sample data.
Run from the repo root: python examples/business-report/build.py [out.pptx]
"""
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parents[1] / "lib"))
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

from deckkit import Deck  # noqa: E402

OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "out.pptx"
d = Deck(HERE / "style.json")
META = "느린여름 로스터리 · 2026년 3분기 사업 검토"
SRC = "출처: 예시 데이터 (가상의 로스터리, POS·구독 관리 시스템 집계, 2026년 9월 30일 기준)"
TOP, BOTTOM = d.content_top, d.content_bottom
BODY_TOP = TOP + 1.0          # report-grid: title / governing message / body
SRC_Y = BOTTOM - 0.3


def governing(s, text):
    """Second tier of the report structure: one muted sentence under the title."""
    d.text(s, d.body_box(0, 10, top=TOP, bottom=TOP + 0.8), text, "governing", color="muted")


def source(s, text=SRC):
    d.text(s, (d.m, SRC_Y, d.col(0, 9)[1], 0.3), text, "caption", color="muted", fit=False)


def figure_rows(s, x, y, w, rows, *, row_h=0.62, accent_row=None):
    """Label / value rows separated by whitespace, one rule above and one below (no rule under every row, D9)."""
    d.rule(s, x, y, w, color="ink")
    for i, (label, value) in enumerate(rows):
        ry = y + 0.14 + i * row_h
        color = "accent" if i == accent_row else "ink"
        d.text(s, (x, ry, w * 0.62, row_h - 0.1), label, color=color, anchor="middle")
        d.text(s, (x + w * 0.62, ry, w * 0.38, row_h - 0.1), value, color=color, align="right", anchor="middle")
    d.rule(s, x, y + 0.28 + len(rows) * row_h, w, color="muted", weight=0.5)


# 1. Cover ---------------------------------------------------------------
x, w = d.col(0, 9)
s = d.slide("3분기, 정기구독이\n매출의 절반을 넘겼다", box=(x, d.H * 0.28, w, 2.2), role="head")
d.text(s, (x, d.H * 0.28 + 2.35, d.col(0, 7)[1], 0.8), "느린여름 로스터리 2026년 3분기 사업 검토", color="muted")
d.text(s, (x, d.H - d.m - 0.3, w, 0.3), "2026.10.14 경영회의 · 운영팀  |  가상의 회사와 예시 데이터입니다",
       "caption", color="accent", fit=False)
d.notes(s, "오늘은 3분기 숫자 하나와 결정 하나를 말씀드리려고 합니다. 이 회사와 숫자는 모두 예시입니다.")

# 2. Big number: the quarter in one figure (scale break) ------------------
s = d.slide("3분기 매출은 1억 8,420만 원으로 1년 전보다 23% 늘었다")
governing(s, "늘어난 3,440만 원보다 구독 증가분(4,260만 원)이 더 크다. 도매와 매장은 줄었다.")
x, w = d.col(0, 7)
d.text(s, (x, BODY_TOP, w, 2.3), "+23%", "head", size=150, color="accent", anchor="bottom")
d.text(s, (x, BODY_TOP + 2.4, w, 0.4), "2026년 3분기 매출, 전년 동기 1억 4,980만 원 대비", "caption", color="muted")
x, w = d.col(8, 4)
d.text(s, (x, BODY_TOP + 0.1, w, 0.35), "채널별 증감 (전년 동기 대비)", "caption", color="muted")
figure_rows(s, x, BODY_TOP + 0.55, w, [("정기구독", "+4,260만 원"), ("카페 도매", "−390만 원"), ("매장 판매", "−430만 원")])
source(s)
d.notes(s, "작년 3분기 1억 4,980만 원에서 올해 1억 8,420만 원이 됐습니다. 그런데 늘어난 건 구독뿐입니다. "
           "도매와 매장은 오히려 줄었고, 구독이 그걸 메우고도 남았습니다. " + SRC)
d.footer(s, META, 2)

# 3. Line chart: subscription overtook wholesale ---------------------------
s = d.slide("구독 매출은 지난해 4분기에 도매를 앞지른 뒤 격차를 넓혔다")
governing(s, "구독자는 1년 사이 610명에서 1,240명으로 늘었고, 도매 거래처는 23곳에서 21곳으로 줄었다.")
x, w = d.col(0, 8)
d.chart(s, (x, BODY_TOP, w, SRC_Y - BODY_TOP - 0.15), "line",
        ["25년 3분기", "4분기", "26년 1분기", "2분기", "3분기"],
        {"정기구독 (백만 원)": [52.0, 63.0, 71.0, 82.0, 94.6], "카페 도매 (백만 원)": [61.0, 60.0, 58.0, 59.0, 57.1]},
        number_format="0")
x, w = d.col(9, 3)
d.text(s, (x, BODY_TOP + 0.3, w, 0.35), "3분기 매출 중 구독 비중", "caption", color="muted")
d.text(s, (x, BODY_TOP + 0.65, w, 1.0), "51.4%", "head", size=54)
d.text(s, (x, BODY_TOP + 1.8, w, 1.4), "1년 전에는 34.7%였다. 같은 기간 매장 판매는 430만 원 줄었다.")
source(s)
d.notes(s, "선이 교차하는 곳이 작년 4분기입니다. 그 뒤로 구독은 분기마다 1,000만 원 안팎씩 늘었고 도매는 제자리입니다. "
           + SRC)
d.footer(s, META, 3)

# 4. Comparison table: margin by channel -----------------------------------
s = d.slide("구독은 매출총이익률이 58%로 세 채널 가운데 가장 높다")
governing(s, "도매보다 19%p 높아서, 같은 매출이라면 이익이 1.5배 남는다.")
rows = [["채널", "3분기 매출", "매출 비중", "매출총이익률", "전년 동기 대비"],
        ["정기구독", "9,460만 원", "51.4%", "58%", "+82%"],
        ["카페 도매", "5,710만 원", "31.0%", "39%", "−6%"],
        ["매장 판매", "3,250만 원", "17.6%", "47%", "−12%"]]
x, w = d.col(0, 9)
row_h = 0.62
gf = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(BODY_TOP), Inches(w), Inches(row_h * len(rows)))
tbl = gf.table
tbl.first_row = tbl.horz_banding = False   # no header fill, no banding: rules only
widths = [2.3, 1.75, 1.55, 1.75]
for c, cw in enumerate(widths + [w - sum(widths)]):
    tbl.columns[c].width = Inches(cw)
for r, row in enumerate(rows):
    tbl.rows[r].height = Inches(row_h)
    for c, value in enumerate(row):
        cell = tbl.cell(r, c)
        cell.fill.background()             # the default table style would tint every cell
        cell.margin_left = cell.margin_right = Inches(0.05)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.text = value
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.RIGHT   # numbers right-aligned
        for run in p.runs:
            run.font.size = Pt(d.style["scale"]["caption" if r == 0 else "body"])
            run.font.color.rgb = d._rgb("muted" if r == 0 else "accent" if r == 1 else "ink")
            run.font._rPr.set("lang", "ko-KR")
d.rule(s, x, BODY_TOP + row_h, w, color="ink")
d.rule(s, x, BODY_TOP + row_h * len(rows), w, color="muted", weight=0.5)
x, w = d.col(9, 3)
d.text(s, (x + 0.3, BODY_TOP + row_h + 0.1, w - 0.3, 2.0),
       "매장 판매는 원두 소매와 음료 매출을 합친 값이다. 도매는 21개 카페 납품분이다.", "caption", color="muted")
source(s)
d.notes(s, "표에서 보실 곳은 첫 줄 하나입니다. 구독이 매출도 가장 크고 이익률도 가장 높습니다. "
           "도매는 물량은 안정적이지만 이익률이 39%에 그칩니다. " + SRC)
d.footer(s, META, 4)

# 5. Bar chart: why subscribers leave --------------------------------------
s = d.slide("3분기 해지 212건 가운데 71건은 배송 지연 때문이었다")
governing(s, "지연을 꼽은 해지는 대부분 화요일 일괄 출고분이었다.")
x, w = d.col(0, 8)
d.chart(s, (x, BODY_TOP, w, SRC_Y - BODY_TOP - 0.15), "bar",
        ["배송 지연", "원두가 남음", "가격 부담", "맛이 취향과 다름", "기타"], {"해지 건수": [71, 54, 38, 27, 22]},
        highlight=0, labels=True)
x, w = d.col(9, 3)
d.text(s, (x, BODY_TOP + 0.1, w, 2.4),
       ["해지 사유 1위, 전체의 33%", "지연 71건 중 60건이\v화요일 출고분"], bullets=True)
source(s, "출처: 예시 데이터 (가상의 로스터리 해지 설문, 2026년 7–9월, n = 212)")
d.notes(s, "해지하신 분께 이유를 하나만 고르게 했습니다. 1위가 배송 지연이고, 그중 60건이 화요일에 나간 물량입니다. "
           "출처: 예시 데이터, 해지 설문 212건.")
d.footer(s, META, 5)

# 6. Diagram: the weekly schedule, now and proposed -------------------------
s = d.slide("로스팅과 출고를 주 2회로 나누면 하루 출고량이 절반이 된다")
governing(s, "지금은 월요일에 볶은 원두 620봉을 화요일 하루에 모두 보낸다.")
days = ["월", "화", "수", "목", "금"]
label_x, label_w = d.col(0, 2)
cells = {"현재": {0: ("로스팅", "muted"), 1: ("출고 620봉", "ink")},
         "제안": {0: ("로스팅", "muted"), 1: ("출고 310봉", "accent"), 3: ("로스팅", "muted"), 4: ("출고 310봉", "accent")}}
y = BODY_TOP + 0.05
for i, day in enumerate(days):
    cx, cw = d.col(2 + 2 * i, 2)
    d.text(s, (cx, y, cw, 0.35), day, "caption", color="muted")
for r, (name, row) in enumerate(cells.items()):
    ry = y + 0.5 + r * 1.05
    d.text(s, (label_x, ry, label_w, 0.85), name, anchor="middle")
    for i, (label, color) in row.items():
        cx, cw = d.col(2 + 2 * i, 2)
        d.rect(s, (cx, ry, cw, 0.85), color)
        d.text(s, (cx + 0.18, ry, cw - 0.3, 0.85), label, color="bg", anchor="middle")
x, w = d.col(2, 10)
d.text(s, (x, y + 2.65, w, 0.5),
       "화요일 출고가 오후 4시 택배 마감을 넘긴 주는 3분기 13주 가운데 9주였다.", color="ink")
source(s, "출처: 예시 데이터 (가상의 로스터리 출고 기록, 2026년 7–9월)")
d.notes(s, "문제는 화요일 하루입니다. 620봉을 하루에 포장하다 보니 13주 중 9주는 택배 마감을 넘겼습니다. "
           "목요일에 한 번 더 볶으면 하루 물량이 310봉으로 줄어듭니다. 출처: 예시 데이터, 출고 기록.")
d.footer(s, META, 6)

# 7. Big number, mirrored: the cost in subscribers --------------------------
s = d.slide("추가 비용 월 180만 원은 구독자 122명을 지키면 메워진다")
governing(s, "3분기에 배송 지연으로 떠난 구독자만 71명이었다. 두 분기면 손익분기를 넘는다.")
x, w = d.col(0, 6)
d.text(s, (x, BODY_TOP + 0.1, w, 0.35), "늘어나는 비용 (월)", "caption", color="muted")
figure_rows(s, x, BODY_TOP + 0.55, w, [("파트타임 포장 인력, 주 2일", "140만 원"), ("택배 픽업 주 2회로 변경", "40만 원"),
                                       ("합계", "180만 원")])
x, w = d.col(7, 5)
d.text(s, (x, BODY_TOP, w, 2.3), "122명", "head", size=130, color="accent", anchor="bottom")
d.text(s, (x, BODY_TOP + 2.4, w, 0.7),
       "월 180만 원 ÷ 구독자 1명당 월 이익 1만 4,750원\v(월 결제 2만 5,430원 × 매출총이익률 58%)", "caption", color="muted")
source(s)
d.notes(s, "비용을 구독자 수로 바꿔 보면 122명입니다. 지연 때문에 떠난 분이 한 분기에 71명이니, "
           "이분들을 지키면 두 분기 안에 비용을 넘습니다. " + SRC)
d.footer(s, META, 7)

# 8. Decision ask -------------------------------------------------------------
s = d.slide("11월부터 석 달간 주 2회 출고를 시범 운영하도록 승인해 주십시오")
governing(s, "오늘 결정할 것은 두 가지이고, 계속할지는 1월 경영회의에서 해지 건수로 판단합니다.")
x, w = d.col(0, 7)
d.text(s, (x, BODY_TOP + 0.1, w, 0.35), "승인 요청", "caption", color="muted")
figure_rows(s, x, BODY_TOP + 0.55, w, [("파트타임 포장 인력 1명 채용", "월 140만 원"),
                                       ("택배 픽업 계약을 주 2회로 변경", "월 40만 원")], row_h=0.8)
x, w = d.col(8, 4)
d.text(s, (x, BODY_TOP + 0.1, w, 0.35), "계속 여부 판단 기준", "caption", color="muted")
d.text(s, (x, BODY_TOP + 0.55, w, 1.6), "11~1월 배송 지연 해지가\v분기 35건 이하", "governing", color="accent")
d.text(s, (x, BODY_TOP + 1.75, w, 1.2), "3분기 71건의 절반이다. 넘으면 2월에 원래 일정으로 돌아간다.")
source(s)
d.notes(s, "승인해 주실 것은 채용 한 명과 택배 계약 변경 하나입니다. 1월에 지연 해지가 35건을 넘으면 되돌리겠습니다. " + SRC)
d.footer(s, META, 8)

d.save(OUT)
