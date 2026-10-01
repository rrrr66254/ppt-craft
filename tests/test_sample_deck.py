import json

import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

import lint
import sample_deck
from deckkit.fonts import find_font_file
from helpers import STYLE

SAMPLE = {
    "title": "캐시 도입 3개월 회고",
    "subtitle": "주문 API 지연을 절반으로 줄인 과정",
    "meta": "2026.10 사내 기술 세미나 · 플랫폼팀",
    "claim": "캐시를 붙이자 p99 지연이 절반으로 줄었다",
    "governing": "읽기 요청 대부분이 캐시에서 끝나면서 DB 부하가 줄었다.",
    "points": ["주문 API p99 820ms에서 410ms로", "읽기 요청의 83%가 캐시 적중", "DB CPU 사용률 61%에서 34%로"],
    "number": {"value": "410ms", "label": "9월 주문 API p99"},
    "chart": {"title": "비용은 9월에 월 38만 원 늘었다", "kind": "column", "categories": ["7월", "8월", "9월"],
              "series": {"비용(만 원)": [120, 131, 158]}, "highlight": 2, "source": "클라우드 청구서",
              "takeaway": "증가분은 캐시 노드 2대 비용이다"},
    "source": "2026년 9월 사내 APM",
}


@pytest.mark.parametrize("cover", ["type", "band"])
def test_builds_three_clean_slides(tmp_path, cover):
    st = sample_deck.resolve_fonts({**STYLE, "cover": cover})
    out = sample_deck.build(st, SAMPLE, tmp_path / "s.pptx")
    prs = Presentation(out)
    assert len(prs.slides) == 3
    assert prs.slides[0].shapes.title.text_frame.text == SAMPLE["title"]
    res = lint.lint(out)
    assert [f for f in res["findings"] if f["severity"] in ("blocker", "major")] == []


def test_band_cover_puts_band_behind_title(tmp_path):
    st = sample_deck.resolve_fonts({**STYLE, "cover": "band"})
    first = Presentation(sample_deck.build(st, SAMPLE, tmp_path / "s.pptx")).slides[0].shapes[0]
    assert not (first.has_text_frame and first.text_frame.text)


def test_image_cover(tmp_path):
    img = tmp_path / "photo.jpg"
    Image.new("RGB", (1600, 900), (90, 120, 150)).save(img)
    st = sample_deck.resolve_fonts({**STYLE, "cover": "image"})
    sample = {**SAMPLE, "image": str(img), "image_focus": [0.4, 0.2, 0.7, 0.8]}
    prs = Presentation(sample_deck.build(st, sample, tmp_path / "s.pptx"))
    assert any(sh.shape_type == MSO_SHAPE_TYPE.PICTURE for sh in prs.slides[0].shapes)


def test_governing_structure_adds_message(tmp_path):
    st = sample_deck.resolve_fonts({**STYLE, "structure": "governing"})
    prs = Presentation(sample_deck.build(st, SAMPLE, tmp_path / "s.pptx"))
    texts = [sh.text_frame.text for sh in prs.slides[1].shapes if sh.has_text_frame]
    assert SAMPLE["governing"] in texts


def test_font_override_falls_back_when_missing(capsys):
    st = sample_deck.resolve_fonts(STYLE, font="NoSuchFont123")
    if any(find_font_file(f) is not None for f in sample_deck.FALLBACKS):
        assert st["font"]["head"] in sample_deck.FALLBACKS
        assert capsys.readouterr().out.count("NoSuchFont123") == 1  # reported only once even if head/body/mono are the same
    else:
        assert st["font"]["head"] == "NoSuchFont123"


def test_cli(tmp_path):
    (tmp_path / "style.json").write_text(json.dumps(STYLE), encoding="utf-8")
    (tmp_path / "sample.json").write_text(json.dumps(SAMPLE, ensure_ascii=False), encoding="utf-8")
    sample_deck.main([str(tmp_path / "style.json"), str(tmp_path / "sample.json"), str(tmp_path / "o.pptx")])
    assert (tmp_path / "o.pptx").exists()


FIGURE = {"caption": "전후 지연", "categories": ["전", "후"], "series": {"지연": [820, 410]}, "highlight": 1}
LAYOUTS = ["split", "statement", "figure"]


def _texts(slide):
    return [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]


def _build(tmp_path, style, sample, capsys):
    """Build the 3 sample slides and check there are no deckkit overflow/broken-word warnings and no lint blockers/majors."""
    out = sample_deck.build(sample_deck.resolve_fonts(style), sample, tmp_path / "s.pptx")
    printed = capsys.readouterr().out
    assert "overflow" not in printed and "broken" not in printed, printed
    bad = [f for f in lint.lint(out)["findings"] if f["severity"] in ("blocker", "major")]
    assert bad == [], bad
    return Presentation(out), printed


@pytest.mark.parametrize("layout", LAYOUTS)
def test_layouts_build_clean(tmp_path, layout, capsys):
    prs, printed = _build(tmp_path, {**STYLE, "layout": layout}, {**SAMPLE, "figure": FIGURE}, capsys)
    assert len(prs.slides) == 3
    assert "[deckkit warning]" not in printed


def test_statement_without_number_falls_back_to_split(tmp_path, capsys):
    sample = {k: v for k, v in SAMPLE.items() if k != "number"}
    prs, printed = _build(tmp_path, {**STYLE, "layout": "statement"}, sample, capsys)
    assert "split" in printed and len(prs.slides) == 3
    assert SAMPLE["points"][2] in _texts(prs.slides[1])[-4:] or any(SAMPLE["points"][2] in t for t in _texts(prs.slides[1]))


def test_figure_without_figure_key_uses_split_content(tmp_path, capsys):
    prs, _ = _build(tmp_path, {**STYLE, "layout": "figure"}, SAMPLE, capsys)
    assert not any(sh.has_chart for sh in prs.slides[1].shapes)  # the same chart is not shown twice
    assert any(sh.has_chart for sh in prs.slides[2].shapes)


ENGLISH = {
    "title": "Three months of caching", "subtitle": "How we halved order API latency", "meta": "Platform team",
    "claim": "Caching cut p99 latency in half", "points": ["p99 fell from 820ms to 410ms", "83% of reads hit the cache"],
    "number": {"value": "410ms", "label": "Sept order API p99"},
    "chart": {"title": "Cost rose by 380k KRW per month", "categories": ["Jul", "Aug", "Sep"],
              "series": {"Cost": [120, 131, 158]}, "highlight": 2, "source": "Cloud invoice"},
    "source": "Internal APM, Sept 2026",
}


@pytest.mark.parametrize("layout", LAYOUTS)
def test_english_sample_uses_english_labels(tmp_path, layout, capsys):
    prs, _ = _build(tmp_path, {**STYLE, "layout": layout}, ENGLISH, capsys)
    all_text = [t for sl in prs.slides for t in _texts(sl)]
    assert "Source: Cloud invoice" in all_text
    assert not any("출처" in t for t in all_text)
    assert prs.slides[1].notes_slide.notes_text_frame.text.endswith("Source: Internal APM, Sept 2026")
    if layout == "figure":
        assert any(t.startswith("Figure 1.") for t in _texts(prs.slides[2]))


def test_missing_source_is_marked(tmp_path, capsys):
    sample = {k: v for k, v in SAMPLE.items() if k != "source"}
    sample["chart"] = {k: v for k, v in SAMPLE["chart"].items() if k != "source"}
    prs, _ = _build(tmp_path, STYLE, sample, capsys)
    assert prs.slides[1].notes_slide.notes_text_frame.text.endswith("[출처 필요]")
    assert "[출처 필요]" in _texts(prs.slides[2])


def test_n_only_when_given(tmp_path, capsys):
    st = {**STYLE, "layout": "figure"}
    prs, _ = _build(tmp_path, st, SAMPLE, capsys)
    assert not any("n =" in t for t in _texts(prs.slides[2]))
    prs, _ = _build(tmp_path, st, {**SAMPLE, "chart": {**SAMPLE["chart"], "n": 12}}, capsys)
    assert any("n = 12" in t for t in _texts(prs.slides[2]))


def test_long_caption_does_not_overlap_source(tmp_path, capsys):
    long_caption = "도입 전후 주문 API p99 지연과 캐시 적중률 변화를 비교한 측정 결과 (최근 세 달, 평일 오전 기준, 이상치 제외)"
    sample = {**SAMPLE, "figure": {**FIGURE, "caption": long_caption, "source": "사내 APM"}}
    prs, _ = _build(tmp_path, {**STYLE, "layout": "figure"}, sample, capsys)
    boxes = {sh.text_frame.text.split(" ")[0]: sh for sh in prs.slides[1].shapes if sh.has_text_frame}
    cap, src = boxes["그림"], boxes["출처:"]
    assert cap.top + cap.height <= src.top + 1000  # EMU; only rounding error is allowed
    assert src.top + src.height <= prs.slide_height - 0.3 * 914400  # ends above the footer line


@pytest.mark.parametrize("layout", LAYOUTS)
def test_long_title_and_six_points(tmp_path, layout, capsys):
    sample = {**SAMPLE, "figure": FIGURE,
              "title": "주문 API 캐시 도입 세 달의 기록과 남은 과제, 그리고 다음 분기 계획",
              "claim": "읽기 요청 대부분을 캐시로 돌리자 주문 API의 p99 지연이 절반으로 줄었다",
              "points": ["p99 820ms에서 410ms로", "읽기 83%가 캐시 적중", "DB CPU 61%에서 34%로",
                         "캐시 노드 2대 추가", "무효화 지연 평균 40ms", "월 비용 38만 원 증가"]}
    prs, _ = _build(tmp_path, {**STYLE, "layout": layout}, sample, capsys)
    assert len(prs.slides) == 3


def test_image_cover_at_4_3_stays_on_panel(tmp_path, capsys):
    img = tmp_path / "photo.jpg"
    Image.new("RGB", (1600, 900), (90, 120, 150)).save(img)
    style = {**STYLE, "cover": "image", "canvas": {"ratio": "4:3", "margin": 0.6}}
    sample = {**SAMPLE, "image": str(img), "image_focus": [0.4, 0.2, 0.7, 0.8]}
    prs, _ = _build(tmp_path, style, sample, capsys)
    panel_right = prs.slide_width * 0.62
    for sh in prs.slides[0].shapes:
        if sh.has_text_frame and sh.text_frame.text:
            assert sh.left + sh.width <= panel_right, sh.text_frame.text
            assert sh.top + sh.height <= prs.slide_height


@pytest.mark.parametrize("bad, match", [({"claim": ""}, "claim"), ({"points": []}, "points"),
                                        ({"chart": {"title": "t"}}, "categories"),
                                        ({"number": {"label": "x"}}, "value")])
def test_invalid_sample_raises_korean_error(tmp_path, bad, match):
    with pytest.raises(ValueError, match=match):
        sample_deck.build(sample_deck.resolve_fonts(STYLE), {**SAMPLE, **bad}, tmp_path / "s.pptx")


def test_points_string_is_normalized(tmp_path, capsys):
    prs, _ = _build(tmp_path, STYLE, {**SAMPLE, "points": "한 줄짜리 근거"}, capsys)
    assert "한 줄짜리 근거" in _texts(prs.slides[1])


def test_number_accepts_plain_string(tmp_path, capsys):
    prs, _ = _build(tmp_path, {**STYLE, "layout": "statement"}, {**SAMPLE, "number": "410ms"}, capsys)
    assert "410ms" in _texts(prs.slides[1])


def test_figure_caption_number_when_layout_fell_back(tmp_path, capsys):
    st = {**STYLE, "layout": "figure"}
    prs, _ = _build(tmp_path, st, SAMPLE, capsys)  # no figure key -> split body
    assert any(t.startswith("그림 1.") for t in _texts(prs.slides[2]))
    prs, _ = _build(tmp_path, st, {**SAMPLE, "figure": FIGURE}, capsys)
    assert any(t.startswith("그림 2.") for t in _texts(prs.slides[2]))
