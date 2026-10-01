import json

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

import lint
from deckkit import Deck
from helpers import STYLE


def make_bad(path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s1 = prs.slides.add_slide(prs.slide_layouts[5])
    s1.shapes.title.text = "Key Takeaways"
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "성장 전략: 혁신적인 미래"
    tb = s2.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1))
    tb.text_frame.text = "🚀 [Insert data] 매출이 300% 늘었다"
    for k in range(3):
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1 + k * 4), Inches(4), Inches(3.5), Inches(2))
        card.fill.gradient()
    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = "데이터 기반 개선"
    tb3 = s3.shapes.add_textbox(Inches(1), Inches(1.8), Inches(10), Inches(1))
    tb3.text_frame.text = "데이터를 통해 확인하고 분석을 통해 검증하며 협업을 통해 개선한다"
    for k, h in enumerate((1.0, 2.0, 3.0)):
        s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2 + k * 1.2), Inches(6.5 - h), Inches(0.8), Inches(h))
    prs.save(path)
    return path


def make_clean(path):
    d = Deck(STYLE)
    s = d.slide("캐시를 붙이자 p99 지연이 절반으로 줄었다")
    d.text(s, d.body_box(0, 7), ["주문 API 기준 820ms에서 410ms로 감소", "읽기 요청의 83%가 캐시에서 처리됨"], bullets=True)
    x, w = d.col(8, 4)
    d.text(s, (x, d.content_top, w, 1.2), "410ms", "head", color="accent")
    d.footer(s, "출처: 2026년 9월 사내 APM", 1)
    s2 = d.slide("비용은 월 38만 원 늘었다")
    d.chart(s2, d.body_box(0, 8), "column", ["7월", "8월", "9월"], {"비용(만 원)": [120, 131, 158]},
            highlight=2, source="클라우드 청구서")
    d.text(s2, d.body_box(9, 3), "증가분은 캐시 노드 2대 비용이다")
    s3 = d.slide("다음 분기에는 쓰기 경로를 줄인다")
    d.text(s3, d.body_box(0, 6), "쓰기 요청을 모아서 보내면 DB 연결 수가 줄어든다.", "governing")
    d.notes(s3, "배치 크기는 실험으로 정한다")
    return d.save(path)


def _style_file(tmp_path):
    p = tmp_path / "style.json"
    p.write_text(json.dumps(STYLE), encoding="utf-8")
    return p


def test_bad_deck_flags_expected_rules(tmp_path):
    res = lint.lint(make_bad(tmp_path / "bad.pptx"))
    rules = {f["rule"] for f in res["findings"]}
    expected = {"S1", "W3", "K4", "I3", "W14", "D3", "C1", "C6", "T12", "T11", "K1", "D5", "T8"}
    assert expected <= rules, expected - rules
    assert res["counts"]["blocker"] >= 3


def test_clean_deck_has_no_blocker_or_major(tmp_path):
    res = lint.lint(make_clean(tmp_path / "clean.pptx"), style=_style_file(tmp_path))
    bad = [f for f in res["findings"] if f["severity"] in ("blocker", "major")]
    assert bad == [], bad


def test_layout_repetition_flagged(tmp_path):
    d = Deck(STYLE)
    for i in range(3):
        s = d.slide(f"제목 {i}")
        d.text(s, d.body_box(0, 7), "본문 내용")
    res = lint.lint(d.save(tmp_path / "rep.pptx"))
    assert any(f["rule"] == "L2" and f["slide"] == 3 for f in res["findings"])


def test_title_underline_flagged(tmp_path):
    d = Deck(STYLE)
    s = d.slide("제목")
    d.rule(s, 0.6, 1.55, 3)
    res = lint.lint(d.save(tmp_path / "u.pptx"))
    assert any(f["rule"] == "L6" for f in res["findings"])


def test_cli_exit_code_and_json(tmp_path):
    out = tmp_path / "lint.json"
    with pytest.raises(SystemExit) as e:
        lint.main([str(make_bad(tmp_path / "bad.pptx")), "--out", str(out)])
    assert e.value.code == 1
    assert json.loads(out.read_text(encoding="utf-8"))["counts"]["blocker"] >= 3


# ---- false-positive regressions ----
def _rules(res):
    return {f["rule"] for f in res["findings"]}


def test_realistic_deck_has_no_false_blocker_or_major(tmp_path):
    d = Deck(STYLE)
    s = d.slide("비용은 월 38만 원 늘었다")
    d.chart(s, d.body_box(0, 8), "column", ["7월", "8월", "9월"], {"비용(만 원)": [120, 131, 158]},
            highlight=2, source="클라우드 청구서")
    d.footer(s, "출처: 2026년 9월 사내 APM", 1)
    s = d.slide("결제 화면에서 이탈이 생긴다")
    d.callout(s, (6, 3.2, 3, 2), "버튼이 스크롤 아래로 밀림", (1.0, d.content_top, 3.0, 0.4))
    s = d.slide("우선순위는 세 가지다")
    d.text(s, d.body_box(0, 8), ["★ 이번 분기 목표", "☞ 다음 단계", "❶ 수집 ➤ 정제"])
    s = d.slide("설정은 환경 변수로 한다")
    d.text(s, d.body_box(0, 8), ["담당자 010-XXXX-XXXX", "Add your API key to .env"])
    s = d.slide("표로 비교한다")
    tbl = s.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1.5)).table
    tbl.cell(0, 0).text = "Plan A"
    tbl.cell(1, 0).text = "Plan B"
    d.notes(s, "표는 말로 설명한다")
    ph = s.notes_slide.notes_placeholder  # a notes slide with no notes placeholder
    ph._element.getparent().remove(ph._element)
    res = lint.lint(d.save(tmp_path / "real.pptx"), style=_style_file(tmp_path))
    bad = [f for f in res["findings"] if f["severity"] in ("blocker", "major")]
    assert bad == [], bad
    assert "W8" not in _rules(res)  # footer and source lines are not a repeated 'keyword: explanation' pattern


def test_table_cell_text_is_linted(tmp_path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tbl = s.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(2)).table
    tbl.cell(0, 0).text = "매출 300% 증가 🚀 [Insert data]"
    prs.save(tmp_path / "t.pptx")
    assert {"W14", "I3", "D3"} <= _rules(lint.lint(tmp_path / "t.pptx"))


def test_keycap_emoji_flagged_but_symbols_are_not():
    keycap = "1" + chr(0xFE0F) + chr(0x20E3)
    assert lint.EMOJI.search(keycap)
    for ch in (0x2605, 0x261E, 0x2776, 0x27A4, 0x2192):  # star, pointing hand, circled digit, arrowhead, arrow
        assert not lint.EMOJI.search(chr(ch))
    for ch in (0x1F680, 0x2705, 0x2728, 0x1F1F0):
        assert lint.EMOJI.search(chr(ch))


def test_keycap_deck_flagged(tmp_path):
    d = Deck(STYLE)
    s = d.slide("단계를 나눈다")
    d.text(s, d.body_box(0, 8), "1" + chr(0xFE0F) + chr(0x20E3) + " 첫 단계")
    assert "I3" in _rules(lint.lint(d.save(tmp_path / "k.pptx")))


def test_placeholder_patterns():
    assert lint.PLACEHOLDER.search("Add your logo here")
    assert lint.PLACEHOLDER.search("텍스트를 입력하세요")
    assert lint.PLACEHOLDER.search("XXX")
    for ok in ("Add your API key to .env", "검색창에 텍스트를 입력하면 뜬다", "010-XXXX-XXXX", "XXXL size"):
        assert not lint.PLACEHOLDER.search(ok), ok


def test_theme_hang_font_counts_as_korean_font(tmp_path):
    from lxml import etree
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    a = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    prs = Presentation()
    part = prs.slide_master.part.part_related_by(RT.THEME)
    root = etree.fromstring(part.blob)
    for f in root.iter(a + "font"):
        if f.get("script") == "Hang":
            f.set("typeface", "Pretendard")
    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "지연이 절반으로 줄었다"
    for r in s.shapes.title.text_frame.paragraphs[0].runs:
        r._r.get_or_add_rPr().set("lang", "ko-KR")
    prs.save(tmp_path / "h.pptx")
    assert "T12" not in _rules(lint.lint(tmp_path / "h.pptx"))


def test_unsourced_number_in_english_word_resource(tmp_path):
    d = Deck(STYLE)
    s = d.slide("GPU resource 사용률이 40% 줄었다")  # 'resource' is not a source marker
    assert "D3" in _rules(lint.lint(d.save(tmp_path / "g.pptx")))


def test_cli_corrupt_file_exits_2_and_removes_stale_output(tmp_path):
    bad = tmp_path / "broken.pptx"
    bad.write_bytes(b"not a zip")
    out = tmp_path / "sub" / "lint.json"
    out.parent.mkdir()
    out.write_text("{}", encoding="utf-8")  # the result of a previous run
    with pytest.raises(SystemExit) as e:
        lint.main([str(bad), "--out", str(out)])
    assert e.value.code == 2
    assert not out.exists()


def test_cli_creates_output_directory(tmp_path):
    out = tmp_path / "new" / "dir" / "lint.json"
    with pytest.raises(SystemExit) as e:
        lint.main([str(make_clean(tmp_path / "c.pptx")), "--out", str(out)])
    assert e.value.code == 0 and out.exists()


# ---- group coordinates, L6/L2 boundaries, layout bullets ----
def _bars(shapes_owner):
    for k, h in enumerate((1.0, 2.0, 3.0)):
        shapes_owner.add_shape(MSO_SHAPE.RECTANGLE, Inches(2 + k * 1.2), Inches(6.5 - h), Inches(0.8), Inches(h))


def test_grouped_fake_bar_chart_flagged(tmp_path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "t"
    _bars(s.shapes.add_group_shape().shapes)  # Ctrl+G: chOff == off, chExt == ext
    prs.save(tmp_path / "g.pptx")
    assert "D5" in _rules(lint.lint(tmp_path / "g.pptx"))


def test_group_child_geometry_is_mapped_to_slide_space(tmp_path):
    from pptx.oxml.ns import qn
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    g = s.shapes.add_group_shape()
    g.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2), Inches(4), Inches(5), Inches(2))
    x = g._element.grpSpPr.find(qn("a:xfrm"))
    for tag, a, b in (("a:off", Inches(1), Inches(5)), ("a:ext", Inches(1), Inches(1)),
                      ("a:chOff", Inches(0), Inches(0)), ("a:chExt", Inches(10), Inches(10))):
        el = x.find(qn(tag))
        el.set("x" if "Off" in tag or tag == "a:off" else "cx", str(a))
        el.set("y" if "Off" in tag or tag == "a:off" else "cy", str(b))
    prs.save(tmp_path / "s.pptx")
    kinds = {sh["kind"]: sh for sh in lint.load(tmp_path / "s.pptx")["slides"][0]["shapes"]}
    child = kinds["rect"]
    assert (round(child["x"], 2), round(child["y"], 2), round(child["w"], 2), round(child["h"], 2)) == (1.2, 5.4, 0.5, 0.2)
    assert kinds["group"]["w"] == pytest.approx(1.0)


def test_thin_accent_rect_under_title_flagged(tmp_path):
    d = Deck(STYLE)
    s = d.slide("제목")
    d.rect(s, (0.6, 1.6, 1.2, 0.06), "accent")
    assert "L6" in _rules(lint.lint(d.save(tmp_path / "ul.pptx")))


def test_footer_only_slides_do_not_trip_layout_repetition(tmp_path):
    style = json.loads(json.dumps(STYLE))
    style["canvas"] = {"margin": 0.8}
    d = Deck(style)
    for i, t in enumerate(["첫 주장은 이렇다", "둘째 주장은 저렇다", "셋째 주장은 그렇다"], 1):
        d.footer(d.slide(t), "2026 사내 세미나", i)
    assert "L2" not in _rules(lint.lint(d.save(tmp_path / "m.pptx")))


def test_section_header_body_is_not_bullets(tmp_path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[2])  # section header: the layout turns bullets off
    s.shapes.title.text = "Section"
    tf = s.placeholders[1].text_frame
    tf.text = "설명 첫 줄입니다"
    tf.add_paragraph().text = "설명 둘째 줄입니다"
    prs.save(tmp_path / "sec.pptx")
    assert [p["bullet"] for p in lint.load(tmp_path / "sec.pptx")["slides"][0]["paras"]] == [False, False]


def test_lint_result_includes_slide_notes(tmp_path):
    d = Deck(STYLE)
    s = d.slide("캐시를 붙이자 지연이 줄었다")
    d.notes(s, "출처: 사내 APM")
    d.slide("노트 없는 장")
    res = lint.lint(d.save(tmp_path / "n.pptx"))
    assert res["notes"] == {"1": "출처: 사내 APM"}


def _w8_rules(tmp_path, lines):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tf = s.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(3)).text_frame
    tf.text = lines[0]
    for line in lines[1:]:
        tf.add_paragraph().text = line
    prs.save(tmp_path / "w8.pptx")
    return _rules(lint.lint(tmp_path / "w8.pptx"))


def test_w8_skips_credit_lines(tmp_path):
    credits = ['Photo: "Squirrel" by Sharp, Wikimedia Commons, CC BY-SA 4.0 (SA)', 'Photo: "Rocks" by Joey, Pexels, Pexels License',
               "AI-generated image (FLUX.1-schnell via Cloudflare Workers AI)", "Icons: Lucide, ISC", '사진: "책상" — 홍길동, Pexels, Pexels License',
               "AI 생성 이미지 (FLUX.1-schnell, Pollinations)"]
    assert "W8" not in _w8_rules(tmp_path, credits)


def test_w8_still_flags_real_keyword_explanation_lists(tmp_path):
    assert "W8" in _w8_rules(tmp_path, ["Speed: much faster than before", "Cost: far lower per request", "Risk: fewer surprises"])
    # a credit line mixed with prose lead-ins does not hide them
    assert "W8" in _w8_rules(tmp_path, ["Photo: x by y, Pexels, Pexels License", "Speed: much faster", "Cost: far lower"])


def test_w8_skips_clock_times(tmp_path):
    assert "W8" not in _w8_rules(tmp_path, ["Started 16:00", "Started 12:00", "Hand-off at 17:00"])

def _w9_rules(tmp_path, lines, bullets):
    d = Deck(STYLE)
    s = d.slide("제목")
    d.text(s, (1, 2, 8, 3), lines, bullets=bullets)
    return _rules(lint.lint(d.save(tmp_path / "w9.pptx")))


UNIFORM = ["서버를 열두 대로 늘렸다", "캐시를 두 겹으로 쌓았다", "로그를 하루씩 모았다"]


def test_w9_flags_uniform_bullets_but_not_plain_paragraphs(tmp_path):
    assert "W9" in _w9_rules(tmp_path, UNIFORM, bullets=True)
    assert "W9" not in _w9_rules(tmp_path, UNIFORM, bullets=False)
