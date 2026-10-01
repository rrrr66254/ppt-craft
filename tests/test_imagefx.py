import colorsys
import json

import pytest
from PIL import Image

import imagefx
from helpers import STYLE
from deckkit.style import load_style


def _sat_mean(im):
    small = im.convert("RGB").resize((8, 8))
    px = [small.getpixel((x, y)) for x in range(8) for y in range(8)]
    return sum(colorsys.rgb_to_hsv(*(c / 255 for c in p))[1] for p in px) / len(px)


def _photo(path, size=(64, 48)):
    im = Image.new("RGB", size)
    im.putdata([(x * 4 % 256, y * 5 % 256, 200) for y in range(size[1]) for x in range(size[0])])
    im.save(path)
    return path


def test_gray_has_equal_channels():
    out = imagefx.gray(Image.new("RGB", (4, 4), (200, 50, 20)))
    r, g, b = out.getpixel((1, 1))
    assert r == g == b


def test_duotone_maps_black_and_white_to_palette():
    im = Image.new("RGB", (2, 1), (0, 0, 0))
    im.putpixel((1, 0), (255, 255, 255))
    out = imagefx.duotone(im, "#112233", "#F0E0D0")
    assert out.getpixel((0, 0)) == pytest.approx((0x11, 0x22, 0x33), abs=2)
    assert out.getpixel((1, 0)) == pytest.approx((0xF0, 0xE0, 0xD0), abs=2)


def test_harmonize_lowers_saturation():
    im = Image.new("RGB", (8, 8), (220, 40, 40))
    assert _sat_mean(imagefx.harmonize(im)) < _sat_mean(im)


def test_tone_darkens_and_desaturates():
    im = Image.new("RGB", (8, 8), (220, 40, 40))
    out = imagefx.tone(im)
    assert _sat_mean(out) < _sat_mean(im) and sum(out.getpixel((0, 0))) < sum(im.getpixel((0, 0)))


@pytest.mark.parametrize("kind", ["paper", "grain", "dots", "grid"])
def test_texture_alpha_capped_and_color_exact(kind):
    t = imagefx.texture(kind, "#336699", 0.06, size=(240, 160))
    assert t.mode == "RGBA" and t.size == (240, 160)
    lo, hi = t.getchannel("A").getextrema()
    assert 0 < hi <= round(0.06 * 255)
    r, g, b, _ = t.split()
    assert r.getextrema() == (0x33, 0x33) and g.getextrema() == (0x66, 0x66) and b.getextrema() == (0x99, 0x99)


def test_texture_is_deterministic():
    a = imagefx.texture("paper", "#000000", 0.05, size=(96, 96), seed=3)
    b = imagefx.texture("paper", "#000000", 0.05, size=(96, 96), seed=3)
    assert a.tobytes() == b.tobytes()


def test_texture_rejects_bad_opacity_and_kind():
    with pytest.raises(ValueError):
        imagefx.texture("paper", "#000000", 0.2)
    with pytest.raises(ValueError):
        imagefx.texture("marble", "#000000", 0.05)


def test_treat_file_writes_new_file_and_keeps_source(tmp_path):
    src = _photo(tmp_path / "p.jpg")
    before = src.read_bytes()
    out = imagefx.treat_file(src, "gray", load_style(STYLE), tmp_path / "treated")
    assert out.name == "p-gray.jpg" and out.exists() and src.read_bytes() == before
    with Image.open(out) as res:
        r, g, b = res.convert("RGB").getpixel((5, 5))
    assert max(r, g, b) - min(r, g, b) <= 2  # gray (JPEG may skew by 1)


def test_treat_file_keeps_alpha_as_png(tmp_path):
    src = tmp_path / "a.png"
    Image.new("RGBA", (16, 16), (200, 30, 30, 90)).save(src)
    out = imagefx.treat_file(src, "harmonize", load_style(STYLE), tmp_path / "t")
    with Image.open(out) as res:
        assert out.suffix == ".png" and res.mode == "RGBA" and res.getpixel((1, 1))[3] == 90


def test_cli_treat_and_duotone_fallback_warning(tmp_path, capsys):
    src = _photo(tmp_path / "p.jpg")
    close = {**STYLE, "color": {**STYLE["color"], "ink": "#777777", "bg": "#888888"}}
    sp = tmp_path / "style.json"
    sp.write_text(json.dumps(close), encoding="utf-8")
    rc = imagefx.main(["treat", str(src), "--mode", "duotone", "--style", str(sp), "--out", str(tmp_path / "o")])
    cap = capsys.readouterr()
    assert rc == 0 and "[imagefx] warning: duotone colors too close; using gray" in cap.err
    with Image.open(tmp_path / "o" / "p-duotone.jpg") as res:
        px = res.convert("RGB").getpixel((5, 5))
    assert max(px) - min(px) <= 2  # the fallback really is gray


def test_cli_texture(tmp_path):
    sp = tmp_path / "style.json"
    sp.write_text(json.dumps(STYLE), encoding="utf-8")
    out = tmp_path / "tex" / "paper.png"
    assert imagefx.main(["texture", "--kind", "paper", "--opacity", "0.06", "--style", str(sp), "--out", str(out)]) == 0
    with Image.open(out) as res:
        assert res.mode == "RGBA" and res.size == (1920, 1080)


def test_cli_exit_codes(tmp_path, capsys):
    sp = tmp_path / "style.json"
    sp.write_text(json.dumps(STYLE), encoding="utf-8")
    assert imagefx.main(["treat", str(tmp_path / "missing.jpg"), "--mode", "gray", "--style", str(sp), "--out", str(tmp_path / "o")]) == 1
    assert imagefx.main(["texture", "--kind", "paper", "--opacity", "0.5", "--style", str(sp), "--out", str(tmp_path / "x.png")]) == 2
    capsys.readouterr()
    with pytest.raises(SystemExit) as e:
        imagefx.main(["treat"])
    assert e.value.code == 2


def _write_style(tmp_path, data=STYLE):
    sp = tmp_path / "style.json"
    sp.write_text(json.dumps(data), encoding="utf-8")
    return str(sp)


def test_cli_output_name_collision_is_skipped(tmp_path, capsys):
    a, b = tmp_path / "a" / "c.jpg", tmp_path / "b" / "c.png"
    for p in (a, b):
        p.parent.mkdir()
        Image.new("RGB", (8, 8), (255, 0, 0)).save(p)
    rc = imagefx.main(["treat", str(a), str(b), "--mode", "gray", "--style", _write_style(tmp_path), "--out", str(tmp_path / "o")])
    err = capsys.readouterr().err
    assert rc == 1 and f"[imagefx] skip {b}: output name collides with" in err
    assert [f.name for f in (tmp_path / "o").iterdir()] == ["c-gray.jpg"]


def test_cli_refuses_to_overwrite_an_input(tmp_path, capsys):
    out = tmp_path / "o"
    out.mkdir()
    src, victim = tmp_path / "p.jpg", out / "p-gray.jpg"
    _photo(src)
    _photo(victim)
    before = victim.read_bytes()
    rc = imagefx.main(["treat", str(src), str(victim), "--mode", "gray", "--style", _write_style(tmp_path), "--out", str(out)])
    assert rc == 1 and victim.read_bytes() == before
    assert "would overwrite an input" in capsys.readouterr().err


def test_palette_transparency_is_kept(tmp_path):
    src = tmp_path / "pal.png"
    pal = Image.new("P", (16, 16), 0)
    pal.putpalette([255, 0, 0, 0, 255, 0] + [0] * 762)
    pal.putpixel((3, 3), 1)
    pal.info["transparency"] = 0
    pal.save(src, transparency=0)
    out = imagefx.treat_file(src, "gray", load_style(STYLE), tmp_path / "o")
    with Image.open(out) as res:
        assert out.suffix == ".png" and res.mode == "RGBA"
        assert res.getpixel((0, 0))[3] == 0 and res.getpixel((3, 3))[3] == 255


def test_duotone_on_dark_style_keeps_dark_shadows():
    im = Image.new("RGB", (2, 1), (0, 0, 0))
    im.putpixel((1, 0), (255, 255, 255))
    out = imagefx.duotone(im, "#F0F0F0", "#101010")  # ink light, bg dark (dark theme)
    assert sum(out.getpixel((0, 0))) < sum(out.getpixel((1, 0)))


def test_duotone_warning_printed_once_per_run(tmp_path, capsys):
    srcs = [_photo(tmp_path / f"p{i}.jpg") for i in range(3)]
    close = {**STYLE, "color": {**STYLE["color"], "ink": "#777777", "bg": "#888888"}}
    imagefx.main(["treat", *map(str, srcs), "--mode", "duotone", "--style", _write_style(tmp_path, close), "--out", str(tmp_path / "o")])
    assert capsys.readouterr().err.count("duotone colors too close") == 1


def test_icc_profile_kept_except_duotone(tmp_path):
    src = tmp_path / "icc.jpg"
    Image.new("RGB", (8, 8), (200, 90, 40)).save(src, icc_profile=b"\0" * 128)
    st = load_style(STYLE)
    with Image.open(imagefx.treat_file(src, "harmonize", st, tmp_path / "o")) as r:
        assert r.info.get("icc_profile") == b"\0" * 128
    with Image.open(imagefx.treat_file(src, "duotone", st, tmp_path / "o")) as r:
        assert "icc_profile" not in r.info


def test_texture_size_follows_canvas_ratio(tmp_path):
    st = {**STYLE, "canvas": {"ratio": "4:3"}}
    assert imagefx.main(["texture", "--kind", "grid", "--style", _write_style(tmp_path, st), "--out", str(tmp_path / "t.png")]) == 0
    with Image.open(tmp_path / "t.png") as r:
        assert r.size == (1920, 1440)


def test_texture_cli_defaults_from_style(tmp_path):
    st = {**STYLE, "imagery": {"texture": {"kind": "dots", "opacity": 0.04}}}
    assert imagefx.main(["texture", "--style", _write_style(tmp_path, st), "--out", str(tmp_path / "t.png")]) == 0
    with Image.open(tmp_path / "t.png") as r:
        assert r.getchannel("A").getextrema()[1] == round(0.04 * 255)


def test_texture_cli_needs_kind_when_style_has_none(tmp_path, capsys):
    assert imagefx.main(["texture", "--style", _write_style(tmp_path), "--out", str(tmp_path / "t.png")]) == 2
    assert "--kind" in capsys.readouterr().err


def test_texture_cli_requires_png_and_wraps_save_errors(tmp_path, capsys):
    sp = _write_style(tmp_path)
    assert imagefx.main(["texture", "--kind", "dots", "--style", sp, "--out", str(tmp_path / "t.jpg")]) == 2
    assert ".png" in capsys.readouterr().err
    blocker = tmp_path / "file"
    blocker.write_text("x")
    assert imagefx.main(["texture", "--kind", "dots", "--style", sp, "--out", str(blocker / "t.png")]) == 1
    assert "[imagefx] error" in capsys.readouterr().err


# ---- compare (Plan 4 Task 8) ----
from pptx import Presentation  # noqa: E402


def _bg_photo(path):
    im = Image.new("RGB", (2400, 1350), (30, 36, 44))
    im.paste((215, 195, 160), (1500, 0, 2400, 1350))
    im.save(path)
    return path


def _fake_render(calls):
    def render(src, out_dir, width=1600, sheet=False, backend=None):
        calls.append((src, out_dir, width))
        out_dir.mkdir(parents=True, exist_ok=True)
        slides = []
        for i in range(len(Presentation(str(src)).slides)):
            p = out_dir / f"slide-{i + 1:02d}.png"
            Image.new("RGB", (960, 540), (i * 30, 90, 120)).save(p)
            slides.append(str(p))
        return {"backend": "fake", "slides": slides}
    return render


def test_compare_builds_seven_slides_and_sheet(tmp_path, monkeypatch, capsys):
    import render
    calls = []
    monkeypatch.setattr(render, "render", _fake_render(calls))
    style = _write_style(tmp_path, {**STYLE, "imagery": {"patterns": ["type-only"]}})  # compare ignores imagery.patterns
    out = tmp_path / "cmp"
    rc = imagefx.main(["compare", str(_bg_photo(tmp_path / "photo.jpg")), "--style", style, "--out", str(out),
                       "--focus", "0.65", "0.1", "1.0", "0.9"])
    cap = capsys.readouterr()
    assert rc == 0, cap.err
    png = out / "compare.png"
    assert cap.out.strip().splitlines()[-1] == str(png) and png.exists()
    prs = Presentation(str(out / "compare.pptx"))
    titles = [s.shapes.title.text_frame.text if s.shapes.title else
              next(sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text)
              for s in prs.slides]
    assert titles == ["none", "harmonize", "gray", "duotone", "bleed-panel", "bleed-scrim", "split"]
    tags = [{sh.name for sh in s.shapes if sh.name.startswith("pc:")} for s in prs.slides]
    assert all("pc:bleed-panel" in t for t in tags[:5])
    assert "pc:bleed-scrim" in tags[5] and "pc:split" in tags[6]
    for mode in ("harmonize", "gray", "duotone"):
        assert (out / "treated" / f"photo-{mode}.jpg").exists()
    assert len(calls) == 1 and calls[0][1] == out / "renders"
    with Image.open(png) as sheet:  # 2 labelled rows of 480px-wide thumbs, 4 columns
        assert sheet.height == 16 + 2 * (270 + 16) and sheet.width > 4 * (480 + 16)


def test_compare_without_renderer_exits_2(tmp_path, monkeypatch, capsys):
    import render

    def missing(*a, **k):
        raise render.RendererMissing(render.INSTALL_HELP)

    monkeypatch.setattr(render, "render", missing)
    rc = imagefx.main(["compare", str(_bg_photo(tmp_path / "photo.jpg")), "--style", _write_style(tmp_path),
                       "--out", str(tmp_path / "cmp")])
    assert rc == 2 and "No renderer found" in capsys.readouterr().err


def test_compare_missing_photo_exits_1(tmp_path, capsys):
    rc = imagefx.main(["compare", str(tmp_path / "nope.jpg"), "--style", _write_style(tmp_path), "--out", str(tmp_path / "c")])
    assert rc == 1 and "[imagefx]" in capsys.readouterr().err
