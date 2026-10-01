import re
import zipfile
from io import BytesIO

import pytest
from PIL import Image
from pptx import Presentation

from deckkit import Deck
from helpers import STYLE


def _chart_xml(pptx):
    z = zipfile.ZipFile(pptx)
    name = next(n for n in z.namelist() if n.startswith("ppt/charts/chart"))
    return z.read(name).decode("utf-8")


def test_chart_native_colored_with_source(tmp_path):
    d = Deck(STYLE)
    s = d.slide("3분기 매출이 가장 컸다")
    d.chart(s, (1, 2, 6, 4), "column", ["1분기", "2분기", "3분기"], {"매출": [10, 12, 15]},
            highlight=2, source="사내 집계")
    out = d.save(tmp_path / "t.pptx")
    xml = _chart_xml(out)
    assert "1D4E89" in xml and "777777" in xml and "gradFill" not in xml
    assert 'typeface="Malgun Gothic"' in xml
    assert '<c:autoTitleDeleted val="1"/>' in xml
    texts = [sh.text_frame.text for sh in Presentation(out).slides[0].shapes if sh.has_text_frame]
    assert "출처: 사내 집계" in texts


def test_chart_labels_show_values_and_hide_value_axis(tmp_path):
    d = Deck(STYLE)
    s = d.slide("3분기 매출이 가장 컸다")
    d.chart(s, (1, 2, 6, 4), "column", ["1분기", "2분기", "3분기"], {"매출": [10, 12, 15]}, labels=True)
    xml = _chart_xml(d.save(tmp_path / "t.pptx"))
    assert '<c:showVal val="1"/>' in xml
    assert '<c:delete val="1"/>' in xml  # valAx deleted


def test_bar_and_column_value_axis_starts_at_zero_line_stays_auto(tmp_path):
    for kind, expect in (("column", True), ("bar", True), ("line", False)):
        for labels in (False, True):
            d = Deck(STYLE)
            d.chart(d.slide("x"), (1, 2, 6, 4), kind, ["a", "b"], {"v": [98, 100]}, labels=labels)
            xml = _chart_xml(d.save(tmp_path / "t.pptx"))
            assert bool(re.search(r'<c:min val="0(\.0)?"/>', xml)) is expect, (kind, labels)  # python-pptx writes 0.0


def test_chart_without_labels_keeps_value_axis(tmp_path):
    d = Deck(STYLE)
    d.chart(d.slide("x"), (1, 2, 6, 4), "column", ["a", "b"], {"v": [1, 2]})
    xml = _chart_xml(d.save(tmp_path / "t.pptx"))
    assert '<c:showVal val="1"/>' not in xml


def test_image_native_crop_keeps_original(tmp_path):
    img = tmp_path / "p.png"
    Image.new("RGB", (400, 200), "red").save(img)
    d = Deck(STYLE)
    s = d.slide()
    pic = d.image(s, (1, 1, 2, 2), img)
    assert abs(pic.crop_left - 0.25) < 1e-4 and abs(pic.crop_right - 0.25) < 1e-4
    shape = Presentation(d.save(tmp_path / "t.pptx")).slides[0].shapes[0]
    assert Image.open(BytesIO(shape.image.blob)).size == (400, 200)


def test_image_must_keep_falls_back_to_contain(tmp_path):
    img = tmp_path / "p.png"
    Image.new("RGB", (400, 200), "red").save(img)
    d = Deck(STYLE)
    pic = d.image(d.slide(), (1, 1, 2, 2), img, must_keep=(0.05, 0.1, 0.95, 0.9))
    assert any("contain" in w for w in d.warnings)
    assert pic.crop_left == 0 and abs(pic.width.inches - 2) < 0.01 and abs(pic.height.inches - 1) < 0.01


def test_low_dpi_warns(tmp_path):
    img = tmp_path / "small.png"
    Image.new("RGB", (100, 100), "blue").save(img)
    d = Deck(STYLE)
    d.image(d.slide(), (1, 1, 4, 4), img)
    assert any("dpi" in w for w in d.warnings)


def test_callout_draws_box_line_label(tmp_path):
    d = Deck(STYLE)
    s = d.slide()
    before = len(s.shapes)
    d.callout(s, (2, 2, 1, 0.5), "여기서 캐시 적중", (4, 1, 2, 0.4))
    assert len(s.shapes) == before + 3
    xml = zipfile.ZipFile(d.save(tmp_path / "t.pptx")).read("ppt/slides/slide1.xml").decode("utf-8")
    assert "<p:cxnSp>" in xml and "여기서 캐시 적중" in xml


def test_line_chart_markers_use_series_color(tmp_path):
    d = Deck(STYLE)
    d.chart(d.slide(), (1, 1, 6, 4), "line", ["a", "b", "c"], {"x": [1, 2, 3]})
    xml = _chart_xml(d.save(tmp_path / "t.pptx"))
    marker = xml[xml.index("<c:marker>"):xml.index("</c:marker>")]
    assert '<c:symbol val="circle"/>' in marker and marker.count("1D4E89") == 2


def test_bar_chart_first_category_on_top(tmp_path):
    d = Deck(STYLE)
    d.chart(d.slide(), (1, 1, 6, 4), "bar", ["a", "b"], {"x": [1, 2]})
    xml = _chart_xml(d.save(tmp_path / "t.pptx"))
    assert '<c:orientation val="maxMin"/>' in xml and '<c:crosses val="max"/>' in xml


def test_highlight_out_of_range_raises():
    d = Deck(STYLE)
    with pytest.raises(ValueError, match="highlight"):
        d.chart(d.slide(), (1, 1, 6, 4), "column", ["a", "b"], {"x": [1, 2]}, highlight=2)


def test_highlight_multi_series_warns():
    d = Deck(STYLE)
    d.chart(d.slide(), (1, 1, 6, 4), "column", ["a", "b"], {"x": [1, 2], "y": [2, 3]}, highlight=0)
    assert any("highlight" in w for w in d.warnings)


def test_callout_left_label_right_aligned(tmp_path):
    d = Deck(STYLE)
    s = d.slide()
    d.callout(s, (6, 2, 1, 0.5), "왼쪽 라벨", (1, 2, 3, 0.4))
    xml = zipfile.ZipFile(d.save(tmp_path / "t.pptx")).read("ppt/slides/slide1.xml").decode("utf-8")
    assert 'algn="r"' in xml


def test_callout_picks_largest_gap_side():
    d = Deck(STYLE)
    s = d.slide()
    d.callout(s, (5, 2, 2, 1), "아래 라벨", (3.4, 5, 1.5, 0.4))  # far below, slightly to the left
    line = next(sh for sh in s.shapes if hasattr(sh, "begin_x"))
    assert abs(line.end_y.inches - 3.0) < 0.01 and abs(line.end_x.inches - 6.0) < 0.01


def test_negative_values_leave_value_axis_auto(tmp_path):
    for kind in ("column", "bar"):
        d = Deck(STYLE)
        d.chart(d.slide("x"), (1, 2, 6, 4), kind, ["a", "b"], {"v": [5, -3]})
        assert not re.search(r'<c:min val="0(\.0)?"/>', _chart_xml(d.save(tmp_path / "t.pptx")))
    d = Deck(STYLE)  # None (gap) values are ignored: the rest is non-negative
    d.chart(d.slide("x"), (1, 2, 6, 4), "column", ["a", "b"], {"v": [5, None]})
    assert re.search(r'<c:min val="0(\.0)?"/>', _chart_xml(d.save(tmp_path / "t.pptx")))
