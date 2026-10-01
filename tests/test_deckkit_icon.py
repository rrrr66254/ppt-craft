import zipfile

import pytest
from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn

import lint
import render
from deckkit import Deck, SVG_EXT_URI, NS_ASVG
from helpers import STYLE

SVG = '<svg xmlns="http://www.w3.org/2000/svg" width="{w}" height="{h}" viewBox="0 0 {w} {h}"><rect width="{w}" height="{h}" fill="{c}"/></svg>'


def _icon_files(tmp_path, w=24, h=24, svg_color="#1D4E89", png_color="#FF0000", png_size=(48, 48)):
    svg, png = tmp_path / "i.svg", tmp_path / "i.png"
    svg.write_text(SVG.format(w=w, h=h, c=svg_color), encoding="utf-8")
    Image.new("RGB", png_size, png_color).save(png)
    return svg, png


def test_icon_xml_relationship_and_content_type(tmp_path):
    svg, png = _icon_files(tmp_path)
    d = Deck(STYLE)
    s = d.slide()
    d.icon(s, (1, 1, 2, 2), svg)  # png_path defaults to i.png next to the svg
    out = d.save(tmp_path / "t.pptx")
    slide = Presentation(out).slides[0]
    blip = slide.shapes[0]._element.find(f".//{qn('a:blip')}")
    ext = blip.find(f"{qn('a:extLst')}/{qn('a:ext')}[@uri='{SVG_EXT_URI}']")
    svg_blip = ext.find(f"{{{NS_ASVG}}}svgBlip")
    rid = svg_blip.get(qn("r:embed"))
    assert slide.part.rels[rid].target_part.partname.endswith(".svg")
    z = zipfile.ZipFile(out)
    names = [n for n in z.namelist() if n.endswith(".svg")]
    assert len(names) == 1 and names[0].startswith("ppt/media/image")
    ct = z.read("[Content_Types].xml").decode("utf-8")
    assert f'PartName="/{names[0]}" ContentType="image/svg+xml"' in ct
    assert z.read(names[0]) == svg.read_bytes()


def test_icon_contain_sizing_non_square_png(tmp_path):
    svg, png = _icon_files(tmp_path, png_size=(200, 100))
    d = Deck(STYLE)
    pic = d.icon(d.slide(), (1, 1, 2, 2), svg, png)
    assert abs(pic.width.inches - 2) < 0.01 and abs(pic.height.inches - 1) < 0.01
    assert abs(pic.top.inches - 1.5) < 0.01


def test_icon_renders_missing_png_with_resvg(tmp_path):
    pytest.importorskip("resvg_py")
    svg = tmp_path / "a.svg"
    svg.write_text(SVG.format(w=24, h=24, c="#1D4E89"), encoding="utf-8")
    d = Deck(STYLE)
    d.icon(d.slide(), (1, 1, 2, 2), svg)
    assert Image.open(tmp_path / "a.png").size == (512, 512)


def test_icon_refuses_unsafe_svg(tmp_path):
    svg, png = _icon_files(tmp_path)
    svg.write_text('<svg xmlns="http://www.w3.org/2000/svg"><image href="http://evil/x.png"/></svg>', encoding="utf-8")
    d = Deck(STYLE)
    with pytest.raises(ValueError, match="external references"):
        d.icon(d.slide(), (1, 1, 2, 2), svg, png)


def test_icon_refuses_non_utf8_svg(tmp_path):
    svg, png = _icon_files(tmp_path)
    svg.write_bytes(b'<svg xmlns="http://www.w3.org/2000/svg"><!-- ' + bytes([255, 254]) + b' --></svg>')
    d = Deck(STYLE)
    with pytest.raises(ValueError, match="UTF-8"):
        d.icon(d.slide(), (1, 1, 2, 2), svg, png)


def test_icon_render_longest_side_512_and_unusable_svg(tmp_path):
    pytest.importorskip("resvg_py")
    wide, tall = tmp_path / "w.svg", tmp_path / "t.svg"
    wide.write_text(SVG.format(w=48, h=24, c="#1D4E89"), encoding="utf-8")
    tall.write_text(SVG.format(w=24, h=48, c="#1D4E89"), encoding="utf-8")
    d = Deck(STYLE)
    s = d.slide()
    d.icon(s, (1, 1, 2, 2), wide)
    d.icon(s, (4, 1, 2, 2), tall)
    assert Image.open(tmp_path / "w.png").size == (512, 256) and Image.open(tmp_path / "t.png").size == (256, 512)
    bad = tmp_path / "bad.svg"
    bad.write_text("<svg", encoding="utf-8")
    with pytest.raises(ValueError, match=r"bad\.svg: could not render"):
        d.icon(s, (1, 1, 2, 2), bad)


def test_icon_warns_when_png_aspect_differs_from_svg(tmp_path):
    svg, png = _icon_files(tmp_path, png_size=(48, 24))  # svg is square
    d = Deck(STYLE)
    d.icon(d.slide(), (1, 1, 2, 2), svg, png)
    assert any("aspect" in w for w in d.warnings)
    d2 = Deck(STYLE)
    svg2, png2 = _icon_files(tmp_path, png_size=(48, 48))
    d2.icon(d2.slide(), (1, 1, 2, 2), svg2, png2)
    assert not d2.warnings
    svg2.write_text('<svg xmlns="http://www.w3.org/2000/svg"><rect/></svg>', encoding="utf-8")  # no size: no warning
    Image.new("RGB", (48, 24), "red").save(png2)
    d3 = Deck(STYLE)
    d3.icon(d3.slide(), (1, 1, 2, 2), svg2, png2)
    assert not d3.warnings


def test_two_icons_sharing_one_png_get_distinct_svg_parts(tmp_path):
    svg1, png = _icon_files(tmp_path)
    svg2 = tmp_path / "j.svg"
    svg2.write_text(SVG.format(w=24, h=24, c="#00FF00"), encoding="utf-8")
    d = Deck(STYLE)
    s = d.slide()
    d.icon(s, (1, 1, 1, 1), svg1, png)
    d.icon(s, (3, 1, 1, 1), svg2, png)
    out = d.save(tmp_path / "t.pptx")
    slide = Presentation(out).slides[0]
    rids = [sh._element.find(f".//{{{NS_ASVG}}}svgBlip").get(qn("r:embed")) for sh in slide.shapes]
    names = [str(slide.part.rels[r].target_part.partname) for r in rids]
    assert len(set(rids)) == 2 and len(set(names)) == 2 and all(n.endswith(".svg") for n in names)
    z = zipfile.ZipFile(out)
    assert {z.read(n.lstrip("/")) for n in names} == {svg1.read_bytes(), svg2.read_bytes()}


def test_icon_is_not_a_rect_for_lint(tmp_path):
    svg, png = _icon_files(tmp_path)

    def findings(with_icon):
        d = Deck(STYLE)
        s = d.slide("제목")
        d.text(s, d.body_box(0, 6), "본문 내용")
        if with_icon:
            d.icon(s, (8, 2, 1, 1), svg, png)
        res = lint.lint(d.save(tmp_path / f"l{with_icon}.pptx"))
        return [(f["rule"], f["slide"]) for f in res["findings"]]

    assert findings(True) == findings(False)


@pytest.mark.skipif(render.detect_backend() != "powerpoint", reason="PowerPoint render only")
def test_powerpoint_shows_svg_not_png_fallback(tmp_path):
    svg, png = _icon_files(tmp_path, png_color="#FF0000")
    d = Deck(STYLE)
    s = d.slide()
    d.icon(s, (1, 1, 3, 3), svg, png)
    src = d.save(tmp_path / "t.pptx")
    res = render.render(src, tmp_path / "renders", width=1000)  # render opens the file in PowerPoint: it must not need repair
    assert len(res["slides"]) == 1
    im = Image.open(res["slides"][0]).convert("RGB")
    cx, cy = round(2.5 / d.W * im.width), round(2.5 / d.H * im.height)
    r, g, b = im.getpixel((cx, cy))
    assert abs(r - 0x1D) < 12 and abs(g - 0x4E) < 12 and abs(b - 0x89) < 12, (r, g, b)
