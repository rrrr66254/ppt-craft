import zipfile

import pytest
from pptx import Presentation

from deckkit import Deck
from helpers import STYLE


def _xml(pptx, name):
    return zipfile.ZipFile(pptx).read(name).decode("utf-8")


def test_theme_neutralized(tmp_path):
    out = Deck(STYLE).save(tmp_path / "t.pptx")
    theme = _xml(out, "ppt/theme/theme1.xml")
    assert "outerShdw" not in theme and "gradFill" not in theme and "bevelT" not in theme
    assert 'typeface="Malgun Gothic"' in theme and "1D4E89" in theme


def test_korean_run_props(tmp_path):
    d = Deck(STYLE)
    s = d.slide("캐시 도입 후 응답 시간이 40% 줄었다")
    d.text(s, d.body_box(0, 7), ["첫 번째 근거입니다", "두 번째 근거"], bullets=True)
    xml = _xml(d.save(tmp_path / "t.pptx"), "ppt/slides/slide1.xml")
    assert xml.count('lang="ko-KR"') >= 3
    assert '<a:ea typeface="Malgun Gothic"/>' in xml
    assert 'buChar char="•"' in xml and ">•" not in xml
    assert 'spc="-' in xml


def test_title_uses_placeholder_at_title_box(tmp_path):
    d = Deck(STYLE)
    d.slide("제목")
    t = Presentation(d.save(tmp_path / "t.pptx")).slides[0].shapes.title
    assert t.text_frame.text == "제목"
    assert abs(t.left.inches - 0.6) < 0.01 and abs(t.top.inches - 0.4) < 0.01


def test_custom_title_box_and_role(tmp_path):
    d = Deck(STYLE)
    d.slide("표지 제목", box=(0.6, 3.0, 9.0, 2.0), role="head")
    t = Presentation(d.save(tmp_path / "t.pptx")).slides[0].shapes.title
    assert abs(t.top.inches - 3.0) < 0.01


def test_no_title_slide(tmp_path):
    d = Deck(STYLE)
    d.slide()
    assert Presentation(d.save(tmp_path / "t.pptx")).slides[0].shapes.title is None


def test_col_grid():
    d = Deck(STYLE)
    x0, w12 = d.col(0, 12)
    assert abs(x0 - 0.6) < 1e-9 and abs(w12 - (d.W - 1.2)) < 1e-6
    x6, w6 = d.col(6, 6)
    assert abs((x6 + w6) - (0.6 + w12)) < 1e-6


def test_rect_has_no_effects_and_no_line(tmp_path):
    d = Deck(STYLE)
    s = d.slide()
    d.rect(s, (1, 1, 2, 1), "accent")
    xml = _xml(d.save(tmp_path / "t.pptx"), "ppt/slides/slide1.xml")
    assert "<a:effectLst/>" in xml and "<a:noFill/></a:ln>" in xml


def test_send_to_back(tmp_path):
    d = Deck(STYLE)
    s = d.slide("제목")
    band = d.rect(s, (0, 0, 13.333, 3), "accent")
    d.send_to_back(s, band)
    first = Presentation(d.save(tmp_path / "t.pptx")).slides[0].shapes[0]
    assert first.shape_id == band.shape_id


def test_overflow_warns():
    d = Deck(STYLE)
    s = d.slide()
    d.text(s, (1, 1, 1, 0.3), "아주 " * 200)
    assert any("overflow" in w for w in d.warnings)


def test_missing_font_warns():
    d = Deck({**STYLE, "font": {"head": "NoSuchFont123", "body": "NoSuchFont123"}})
    assert any("NoSuchFont123" in w for w in d.warnings)


def test_notes_and_footer(tmp_path):
    d = Deck(STYLE)
    s = d.slide("t")
    d.notes(s, "말할 내용")
    d.footer(s, "2026.10 세미나", 3)
    p = Presentation(d.save(tmp_path / "t.pptx")).slides[0]
    assert p.notes_slide.notes_text_frame.text == "말할 내용"
    assert any(sh.has_text_frame and sh.text_frame.text == "3" for sh in p.shapes)


def test_bad_color_name_raises():
    d = Deck(STYLE)
    with pytest.raises(ValueError, match="color"):
        d.rect(d.slide(), (1, 1, 1, 1), "purple")


def test_notes_master_theme_neutralized(tmp_path):
    d = Deck(STYLE)
    d.notes(d.slide("t"), "메모")
    d.notes(d.slide("u"), "둘째")  # a second call is fine too (idempotent)
    z = zipfile.ZipFile(d.save(tmp_path / "t.pptx"))
    themes = [n for n in z.namelist() if n.startswith("ppt/theme/theme")]
    assert len(themes) == 2
    for n in themes:
        xml = z.read(n).decode("utf-8")
        assert "outerShdw" not in xml and "gradFill" not in xml and "bevelT" not in xml


def test_no_default_office_accents(tmp_path):
    theme = _xml(Deck(STYLE).save(tmp_path / "t.pptx"), "ppt/theme/theme1.xml")
    for default in ("9BBB59", "8064A2", "4BACC6", "F79646"):
        assert default not in theme


@pytest.mark.parametrize("kwargs", [{"role": "nope"}, {"align": "justify"}, {"anchor": "side"}])
def test_bad_option_raises_before_shape(kwargs):
    d = Deck(STYLE)
    s = d.slide()
    before = len(s.shapes)
    with pytest.raises(ValueError):
        d.text(s, (1, 1, 2, 1), "x", **kwargs)
    assert len(s.shapes) == before


def test_bad_slide_role_leaves_no_slide():
    d = Deck(STYLE)
    with pytest.raises(ValueError, match="role"):
        d.slide("t", role="nope")
    assert len(d.prs.slides) == 0


def test_slide_title_newline_becomes_paragraphs(tmp_path):
    d = Deck(STYLE)
    d.slide("첫 줄\n둘째 줄")
    t = Presentation(d.save(tmp_path / "t.pptx")).slides[0].shapes.title
    assert [p.text for p in t.text_frame.paragraphs] == ["첫 줄", "둘째 줄"]
    assert t.text_frame.paragraphs[0].space_after in (None, 0)


def test_footer_box_height(tmp_path):
    d = Deck(STYLE)
    s = d.slide()
    d.footer(s, "행사", 1)
    assert all(abs(sh.height.inches - 0.3) < 0.01 for sh in s.shapes)


def test_save_survives_narrow_console(tmp_path, monkeypatch):
    import builtins

    def strict_print(msg="", *a, **k):
        msg.encode("cp949")  # raises UnicodeEncodeError if there is an unencodable character
        return None

    d = Deck({**STYLE, "font": {"head": "NoSuchFont☃", "body": "NoSuchFont☃"}})
    monkeypatch.setattr(builtins, "print", strict_print)
    d.save(tmp_path / "t.pptx")  # passes if it finishes without an exception


def test_text_size_override(tmp_path):
    d = Deck(STYLE)
    s = d.slide()
    d.text(s, (1, 1, 6, 3), "410ms", "head", size=96)
    with pytest.raises(ValueError, match="size"):
        d.text(s, (1, 1, 6, 2), "x", size=0)
    sh = Presentation(d.save(tmp_path / "t.pptx")).slides[0].shapes[0]
    assert sh.text_frame.paragraphs[0].runs[0].font.size.pt == 96


def test_vertical_tab_is_a_line_break_inside_one_paragraph(tmp_path):
    from pptx.oxml.ns import qn
    d = Deck(STYLE)
    s = d.slide("제목")
    sh = d.text(s, (1, 2, 8, 3), ["첫 문장은 여기서\v끊고 이어집니다", "다음 문단"])
    paras = sh.text_frame.paragraphs
    assert len(paras) == 2
    assert paras[0].text == "첫 문장은 여기서\v끊고 이어집니다"
    assert [r.text for r in paras[0].runs] == ["첫 문장은 여기서", "끊고 이어집니다"]
    br = paras[0]._p.findall(qn("a:br"))
    assert len(br) == 1 and br[0].find(qn("a:rPr")).get("sz") == str(round(paras[0].runs[0].font.size.pt * 100))
