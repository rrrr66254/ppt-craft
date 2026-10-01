import random

import pytest
from PIL import Image
from pptx.util import Emu

import render
from deckkit import Deck, legibility as L
from deckkit.patterns import suggest
from deckkit.crop import cover_crop
from helpers import STYLE

EMU = 914400


def _photo(tmp_path, name="p.jpg", size=(2400, 1350), color=(30, 30, 30)):
    p = tmp_path / name
    Image.new("RGB", size, color).save(p)
    return p


def _two_tone(tmp_path, dark=35, light=225, cut=0.55, size=(1600, 900), name="n.png"):
    """Dark left part, light right part: no single text color passes over the 7-column text box without a scrim."""
    img = Image.new("RGB", size, (light,) * 3)
    img.paste((dark,) * 3, (0, 0, round(size[0] * cut), size[1]))
    p = tmp_path / name
    img.save(p)
    return p


def _on_grid(d, x):
    return any(abs(x - d.col(i, n)[0]) < 1e-6 for n in range(1, 13) for i in range(n))


def _inside(d, box):
    x, y, w, h = box
    return x >= 0 and y >= 0 and x + w <= d.W + 1e-6 and y + h <= d.H + 1e-6 and w > 0 and h > 0


def _inches(shape):
    return tuple(v / EMU for v in (shape.left, shape.top, shape.width, shape.height))


def _names(slide):
    return [sh.name for sh in slide.shapes]


def _check_box(d, box):
    assert _inside(d, box) and _on_grid(d, box[0])


@pytest.fixture
def d():
    return Deck(STYLE)


# ---- bleed_panel ----
@pytest.mark.parametrize("side,focus", [("left", None), ("right", None), (None, (0.0, 0.0, 0.3, 1.0))])
def test_bleed_panel_box_tags_and_zorder(d, tmp_path, side, focus):
    s, box = d.pattern("bleed-panel", "Panel title", _photo(tmp_path), side=side, focus=focus)
    _check_box(d, box)
    start = 12 - 5 if (side or "right") == "right" else 0
    assert (box[0], box[2]) == pytest.approx(d.col(start, 5))
    shapes = list(s.shapes)
    assert [sh.name for sh in shapes][:2] == ["pc:bleed-panel", "pc:panel"]
    assert shapes[2].has_text_frame and shapes[2].text_frame.text == "Panel title"  # background, panel, then text
    px, py, pw, ph = _inches(shapes[1])
    assert py == 0 and ph == pytest.approx(d.H, abs=1e-3)
    on_right = (side or "right") == "right"
    assert (px > 0) == on_right
    assert (box[0] > d.W / 2) == on_right
    assert px - 1e-6 <= box[0] <= px + pw  # text sits inside the panel


def test_bleed_panel_title_inside_panel(d, tmp_path):
    s, box = d.pattern("bleed-panel", "T", _photo(tmp_path), side="left")
    _, panel, title = list(s.shapes)[:3]
    assert panel.left <= title.left and title.left + title.width <= panel.left + panel.width
    assert (title.top + title.height) / EMU <= box[1]


# ---- bleed_scrim ----
def test_bleed_scrim_dark_photo_no_scrim_light_title(d, tmp_path):
    s, box = d.pattern("bleed-scrim", "Quiet dark", _photo(tmp_path), side="left")
    shapes = list(s.shapes)
    assert [sh.name for sh in shapes][0] == "pc:bleed-scrim" and len(shapes) == 2  # background + title only
    assert box is None and d.text_color == "bg"
    assert str(shapes[1].text_frame.paragraphs[0].runs[0].font.color.rgb) == "FFFFFF"


def test_bleed_scrim_gray_photo_adds_scrim_with_planned_alpha(d, tmp_path):
    img = _two_tone(tmp_path)
    s, _ = d.pattern("bleed-scrim", "Gray photo", img, side="left")
    names = _names(s)
    assert names[:2] == ["pc:bleed-scrim", "pc:scrim"] and len(names) == 3  # bg, scrim, title
    alpha = int(list(s.shapes)[1]._element.spPr.xpath(".//a:alpha/@val")[0]) / 100000
    x, w = d.col(0, 7)
    _, ty, _, th = d.style["canvas"]["title_box"]
    with Image.open(img) as im:
        crop = cover_crop(im.width, im.height, d.W, d.H)
        px = L.region_pixels(im, L.slide_box_to_image_px((x, ty, w, th), (0, 0, d.W, d.H), im.size, crop))
    plan = L.plan(px, STYLE["color"]["ink"], STYLE["color"]["bg"], d.style["scale"]["title"])
    assert not plan["panel"] and alpha == pytest.approx(max(plan["alpha"], 0.35))  # spec 2 floor
    assert d.text_color == plan["text"]
    scrim_rgb = list(s.shapes)[1]._element.spPr.xpath(".//a:srgbClr/@val")[0]
    opposite = "ink" if plan["text"] == "bg" else "bg"
    assert "#" + scrim_rgb == STYLE["color"][opposite].upper()  # scrim colour is opposite the text colour


@pytest.mark.parametrize("side", ["left", "right"])
def test_bleed_scrim_is_a_full_height_band_from_the_text_side_edge(d, tmp_path, side):
    s, _ = d.pattern("bleed-scrim", "Band", _two_tone(tmp_path), side=side)
    scrim = next(sh for sh in s.shapes if sh.name == "pc:scrim")
    x, y, w, h = _inches(scrim)
    tx, tw = d.col(0 if side == "left" else 5, 7)
    assert (y, h) == pytest.approx((0, d.H), abs=1e-3)
    if side == "left":
        assert (x, x + w) == pytest.approx((0, tx + tw + 0.3), abs=1e-3)
    else:
        assert (x, x + w) == pytest.approx((tx - 0.3, d.W), abs=1e-3)


def test_bleed_scrim_falls_back_to_panel_when_palette_cannot_reach_contrast(tmp_path):
    low = {**STYLE, "color": {**STYLE["color"], "ink": "#777777", "bg": "#888888"}}
    d = Deck(low)
    s, box = d.pattern("bleed-scrim", "Mid gray", _photo(tmp_path, color=(128, 128, 128)), side="left")
    assert _names(s)[:2] == ["pc:bleed-panel", "pc:panel"]
    assert any("fails contrast" in w for w in d.warnings)
    _check_box(d, box)


def test_bleed_scrim_many_words_uses_panel(d, tmp_path):
    s, box = d.pattern("bleed-scrim", "Wordy", _photo(tmp_path), words=20)
    assert "pc:panel" in _names(s) and box is not None
    assert d._bleeds == 1 and len(d.prs.slides) == 1  # no stray slide or background from the abandoned scrim
    assert any("20 words" in w for w in d.warnings)


def test_bleed_scrim_unplannable_region_falls_back_to_panel(d, tmp_path, monkeypatch):
    def boom(*a, **k):
        raise ValueError("empty region")
    monkeypatch.setattr(L, "plan", boom)
    s, _ = d.pattern("bleed-scrim", "x", _photo(tmp_path))
    assert "pc:panel" in _names(s) and any("could not be analyzed" in w for w in d.warnings)


def test_bleed_scrim_unfittable_must_keep_has_its_own_warning(d, tmp_path):
    img = _photo(tmp_path, size=(3000, 1000))
    s, _ = d.pattern("bleed-scrim", "x", img, must_keep=(0.0, 0.0, 1.0, 1.0))
    assert "pc:bleed-panel" in _names(s)
    assert any("must_keep" in w and "using bleed-panel" in w for w in d.warnings)
    assert not any("fails contrast" in w for w in d.warnings)


def test_bleed_scrim_with_body_returns_grid_box(d, tmp_path):
    s, box = d.pattern("bleed-scrim", "Short body", _photo(tmp_path), words=8, side="right")
    _check_box(d, box)
    assert (box[0], box[2]) == pytest.approx(d.col(5, 7))
    assert box[1] == pytest.approx(d.content_top)


def test_bleed_scrim_transparent_pixels_read_as_slide_bg(d, tmp_path):
    p = tmp_path / "clear.png"
    Image.new("RGBA", (2400, 1350), (0, 0, 0, 0)).save(p)
    d.pattern("bleed-scrim", "Clear", p, side="left")
    assert d.text_color == "ink"  # read as black, the title would turn light


def test_bleed_scrim_samples_raw_pixels_not_exif_rotated(d, tmp_path):
    # PowerPoint draws the raw orientation, so the title region must be read from the raw pixels.
    # Raw 1600x900: dark top half, light bottom half; the title sits at the top -> light text, no scrim.
    # (An EXIF-rotated read would put the dark part at the right and give a mixed region.)
    img = Image.new("RGB", (1600, 900), (225, 225, 225))
    img.paste((30, 30, 30), (0, 0, 1600, 450))
    exif = img.getexif()
    exif[0x0112] = 6
    p = tmp_path / "rot.jpg"
    img.save(p, exif=exif)
    s, _ = d.pattern("bleed-scrim", "Raw orientation", p, side="left")
    assert d.text_color == "bg"
    assert _names(s)[0] == "pc:bleed-scrim" and len(_names(s)) == 2  # background + title, no scrim, no panel
    assert any("EXIF" in w for w in d.warnings)  # the existing d.image warning is kept


# ---- split ----
@pytest.mark.parametrize("focus,text_right", [((0.0, 0.0, 0.3, 1.0), True), ((0.7, 0.0, 1.0, 1.0), False), (None, False)])
def test_split_is_a_half_bleed_opposite_the_focus(d, tmp_path, focus, text_right):
    s, box = d.pattern("split", "Split", _photo(tmp_path), focus=focus)
    _check_box(d, box)
    assert (box[0], box[2]) == pytest.approx(d.col(5, 7) if text_right else d.col(0, 7))
    pic = next(sh for sh in s.shapes if sh.name == "pc:split")
    x, y, w, h = _inches(pic)
    assert (y, h) == pytest.approx((0, d.H), abs=1e-3)  # full slide height
    if text_right:   # image left: from the slide edge to the image columns' far grid edge
        assert (x, x + w) == pytest.approx((0, sum(d.col(0, 5))), abs=1e-3)
    else:            # image right: from the image columns' grid edge to the slide edge
        assert (x, x + w) == pytest.approx((d.col(7, 5)[0], d.W), abs=1e-3)
    assert _on_grid(d, x) or x == 0


def test_split_title_is_confined_to_the_text_columns(d, tmp_path):
    long_title = "사진 옆에 놓인 아주 긴 제목은 사진 아래로 흘러 들어가면 안 된다"
    for side in ("left", "right"):
        s, box = d.pattern("split", long_title, _photo(tmp_path), side=side)
        title = s.shapes.title
        assert (title.left / EMU, title.width / EMU) == pytest.approx((box[0], box[2]), abs=1e-3)


def test_split_ratio_and_tall_image_gets_narrow_part(d, tmp_path):
    s, box = d.pattern("split", "Wide", _photo(tmp_path), ratio=(5, 7), side="left")
    pic = next(sh for sh in s.shapes if sh.name == "pc:split")
    assert pic.left / EMU == pytest.approx(d.col(7, 5)[0])
    s, box = d.pattern("split", "Tall", _photo(tmp_path, "t.jpg", (1000, 1600)), ratio=(8, 4), side="left")
    pic = next(sh for sh in s.shapes if sh.name == "pc:split")
    assert pic.left / EMU == pytest.approx(d.col(8, 4)[0])  # tall photo: the 4-column part
    assert box[2] == pytest.approx(d.col(0, 8)[1])


def test_split_cover_crop_follows_focus_in_the_half_bleed_box(d, tmp_path):
    s, _ = d.pattern("split", "Crop", _photo(tmp_path, size=(3000, 1000)), focus=(0.0, 0.0, 0.2, 1.0), side="right")
    pic = next(sh for sh in s.shapes if sh.name == "pc:split")
    x, y, w, h = _inches(pic)
    assert pic.crop_left == pytest.approx(0) and pic.crop_right > 0.3
    assert (w / h) == pytest.approx(3000 * (1 - pic.crop_right) / 1000, rel=1e-2)  # undistorted


def test_split_footer_text_stays_clear_of_the_photo_on_the_text_side(d, tmp_path):
    s, box = d.pattern("split", "Footer", _photo(tmp_path), side="left")
    d.footer(s, "Event")
    pic = next(sh for sh in s.shapes if sh.name == "pc:split")
    left = next(sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text == "Event")
    assert left.left + left.width * 0.5 < pic.left  # the short footer text itself sits on the text side


@pytest.mark.parametrize("side", ["left", "right"])
def test_split_footer_boxes_stay_off_the_photo(d, tmp_path, side):
    s, _ = d.pattern("split", "Footer", _photo(tmp_path), side=side)
    d.footer(s, "Event", 4)
    pic = next(sh for sh in s.shapes if sh.name == "pc:split")
    for f in (sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text in ("Event", "4")):
        assert f.left + f.width <= pic.left or f.left >= pic.left + pic.width


@pytest.mark.parametrize("side", ["left", "right"])
def test_split_contain_aligns_to_grid_edge_of_its_side(d, tmp_path, side):
    s, box = d.pattern("split", "Contain", _photo(tmp_path, "t.jpg", (1000, 1600)), side=side, fit="contain", ratio=(6, 6))
    pic = next(sh for sh in s.shapes if sh.name == "pc:split")
    x, y, w, h = _inches(pic)
    ix, iw = d.col(6, 6) if side == "left" else d.col(0, 6)
    assert w < iw - 0.1  # a tall photo does not fill the slot
    if side == "left":   # image on the right: flush with the right grid edge
        assert x + w == pytest.approx(ix + iw, abs=1e-3)
    else:
        assert x == pytest.approx(ix, abs=1e-3)


def test_split_bad_ratio(d, tmp_path):
    with pytest.raises(ValueError, match="ratio"):
        d.pattern("split", "x", _photo(tmp_path), ratio=(5, 5))


# ---- inset ----
def test_inset_caption_below_image_and_box(d, tmp_path):
    s, box = d.pattern("inset", "Inset", _photo(tmp_path), caption="Photo: Han, 2026, Seoul")
    pic = next(sh for sh in s.shapes if sh.name == "pc:inset")
    cap = next(sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.startswith("Photo:"))
    assert cap.top >= pic.top + pic.height and cap.left == pic.left
    assert _on_grid(d, pic.left / EMU)
    _check_box(d, box)
    assert (box[0], box[2]) == pytest.approx(d.col(8, 4))


def test_inset_without_caption_raises(d, tmp_path):
    with pytest.raises(ValueError, match="I13"):
        d.pattern("inset", "x", _photo(tmp_path), caption="  ")
    with pytest.raises(TypeError):
        d.pattern("inset", "x", _photo(tmp_path))


def test_inset_full_span_has_no_content_box(d, tmp_path):
    _, box = d.pattern("inset", "x", _photo(tmp_path), caption="c", span=12)
    assert box is None


# ---- strip ----
@pytest.mark.parametrize("position", ["bottom", "top"])
def test_strip_band_and_box(d, tmp_path, position):
    s, box = d.pattern("strip", "Strip", _photo(tmp_path, size=(4000, 1000)), position=position, height_ratio=0.35)
    pic = next(sh for sh in s.shapes if sh.name == "pc:strip")
    x, y, w, h = _inches(pic)
    assert w == pytest.approx(d.W, abs=1e-3) and h == pytest.approx(d.H * 0.35, abs=1e-3)
    assert (y == 0) == (position == "top")
    _check_box(d, box)
    if position == "top":
        assert box[1] >= y + h - 1e-6
    else:
        assert box[1] + box[3] <= y + 1e-6


def test_bottom_strip_leaves_the_footer_zone_free(d, tmp_path):
    s, _ = d.pattern("strip", "Strip", _photo(tmp_path, size=(4000, 1000)), height_ratio=0.4)
    d.footer(s, "Event", 3)
    pic = next(sh for sh in s.shapes if sh.name == "pc:strip")
    footers = [sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text in ("Event", "3")]
    assert len(footers) == 2
    assert all(f.top >= pic.top + pic.height for f in footers)


def test_strip_too_tall_raises(d, tmp_path):
    with pytest.raises(ValueError, match="height_ratio"):
        d.pattern("strip", "x", _photo(tmp_path), height_ratio=0.5)


# ---- gallery ----
@pytest.mark.parametrize("n", [2, 3, 4])
def test_gallery_n_cells_one_row_same_crop(d, tmp_path, n):
    paths = [_photo(tmp_path, f"g{i}.jpg") for i in range(n)]
    s, box = d.pattern("gallery", "Gallery", paths, captions=[f"c{i}" for i in range(n)])
    pics = [sh for sh in s.shapes if sh.name == "pc:gallery"]
    assert len(pics) == n and box is None
    assert len({(p.top, p.height, p.width) for p in pics}) == 1
    assert all(_on_grid(d, p.left / EMU) and _inside(d, _inches(p)) for p in pics)
    assert [p.left for p in pics] == sorted(p.left for p in pics)


def _gallery(d, tmp_path, n, hero, captions=True):
    paths = [_photo(tmp_path, f"g{i}.jpg") for i in range(n)]
    caps = [f"c{i}" for i in range(n)] if captions else None
    s, _ = d.pattern("gallery", "Gallery", paths, hero=hero, captions=caps)
    pics = [sh for sh in s.shapes if sh.name == "pc:gallery"]
    texts = {sh.text_frame.text: sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.startswith("c")}
    return pics, [texts[c] for c in (caps or [])]


def _bottom(shape):
    return (shape.top + shape.height) / EMU


@pytest.mark.parametrize("n,hero", [(2, 0), (3, 0), (3, 1), (3, 2)])
def test_gallery_hero_stacks_the_others_beside_it_to_the_same_height(d, tmp_path, n, hero):
    pics, caps = _gallery(d, tmp_path, n, hero)
    h = pics[hero]
    others = [pics[i] for i in range(n) if i != hero]
    top = d.content_top
    assert h.top / EMU == pytest.approx(top) and others[0].top / EMU == pytest.approx(top)
    # one column of equal cells beside the hero, in list order
    assert len({o.left for o in others}) == 1 and len({(o.width, o.height) for o in others}) == 1
    assert [o.top for o in others] == sorted(o.top for o in others)
    left_hero = hero != n - 1
    assert (h.left < others[0].left) == left_hero
    assert h.left / EMU == pytest.approx(d.col(0, 8)[0] if left_hero else d.col(4, 8)[0])
    assert others[0].left / EMU == pytest.approx(d.col(8, 4)[0] if left_hero else d.col(0, 4)[0])
    assert h.width > 1.9 * others[0].width
    # captions under every cell, and the stack ends where the hero (with its caption) ends
    for pic, cap in zip(pics, caps):
        assert cap.top >= pic.top + pic.height and cap.left == pic.left
        assert _bottom(cap) <= d.content_bottom + 1e-3
    last_other = max(i for i in range(n) if i != hero)
    assert _bottom(caps[hero]) == pytest.approx(_bottom(caps[last_other]), abs=1e-3)
    assert _bottom(caps[hero]) == pytest.approx(d.content_bottom, abs=0.01)


def test_gallery_hero_without_captions_fills_the_height(d, tmp_path):
    pics, _ = _gallery(d, tmp_path, 3, 0, captions=False)
    others = pics[1:]
    assert len({o.left for o in others}) == 1 and _bottom(others[-1]) == pytest.approx(_bottom(pics[0]), abs=1e-3)
    assert _bottom(pics[0]) == pytest.approx(d.content_bottom, abs=1e-3)


def test_gallery_hero_of_four_has_no_room_and_falls_back_to_equal_cells(d, tmp_path):
    pics, _ = _gallery(d, tmp_path, 4, 0)
    assert len({(p.top, p.width, p.height) for p in pics}) == 1  # the equal row
    assert any("hero" in w and "equal" in w for w in d.warnings)


@pytest.mark.parametrize("n", [2, 3, 4])
def test_gallery_without_hero_is_one_row_of_equal_cells(d, tmp_path, n):
    pics, _ = _gallery(d, tmp_path, n, None)
    assert len({(p.top, p.width, p.height) for p in pics}) == 1
    assert [p.left for p in pics] == sorted(p.left for p in pics)


def test_gallery_count_limits(d, tmp_path):
    for n in (0, 1, 5):
        with pytest.raises(ValueError, match="2 to 4"):
            d.pattern("gallery", "x", [_photo(tmp_path, f"{i}.jpg") for i in range(n)])


def test_gallery_without_captions_warns(d, tmp_path):
    d.pattern("gallery", "x", [_photo(tmp_path, "a.jpg"), _photo(tmp_path, "b.jpg")])
    assert any("I13" in w for w in d.warnings)


# ---- shared rules ----
KEEP, FOCUS = (0.4, 0.3, 0.6, 0.7), (0.3, 0.2, 0.7, 0.8)


@pytest.mark.parametrize("name", ["bleed-panel", "bleed-scrim", "split", "inset", "strip", "gallery"])
def test_every_helper_accepts_the_common_kwargs(d, tmp_path, name):
    img = _photo(tmp_path)
    extra = {"caption": "c"} if name == "inset" else {}
    target = [img, img, img] if name == "gallery" else img
    s, box = d.pattern(name, "T", target, focus=FOCUS, must_keep=KEEP, words=3, **extra)
    assert len(s.shapes) >= 2 and (box is None or _inside(d, box))


def test_gallery_per_image_focus_and_must_keep(d, tmp_path):
    imgs = [_photo(tmp_path, f"g{i}.jpg", size=(3000, 1000)) for i in range(3)]
    s, _ = d.pattern("gallery", "G", imgs, captions=["a", "b", "c"],
                     focus=[(0.0, 0.0, 0.3, 1.0), None, (0.7, 0.0, 1.0, 1.0)], must_keep=[None, KEEP, None])
    pics = [sh for sh in s.shapes if sh.name == "pc:gallery"]
    assert pics[0].crop_left < pics[2].crop_left  # per-image focus reached each crop
    with pytest.raises(ValueError, match="one entry per image"):
        d.pattern("gallery", "G", imgs, captions=["a", "b", "c"], focus=[None, None])


def test_slide_resets_text_color(d, tmp_path):
    d.pattern("bleed-scrim", "dark", _photo(tmp_path))
    assert d.text_color == "bg"
    d.slide("plain")
    assert d.text_color == "ink"


def test_only_bleed_patterns_count_toward_max_bleed(tmp_path):
    d = Deck({**STYLE, "imagery": {"max_bleed": 1}})
    img = _photo(tmp_path)
    for _ in range(3):
        d.pattern("split", "s", img)
        d.pattern("strip", "s", img)
        d.pattern("inset", "s", img, caption="c")
    assert d._bleeds == 0 and not any("L24" in w for w in d.warnings)
    d.pattern("bleed-panel", "b", img)
    assert d._bleeds == 1 and not any("L24" in w for w in d.warnings)

def test_disallowed_pattern_raises_with_allowed_list(tmp_path):
    d = Deck({**STYLE, "imagery": {"patterns": ["split", "type-only"]}})
    with pytest.raises(ValueError, match=r"not allowed by style.imagery.patterns \(\['split', 'type-only'\]\)"):
        d.pattern("bleed-panel", "x", _photo(tmp_path))


def test_unknown_pattern_name(d, tmp_path):
    with pytest.raises(ValueError, match="pattern must be one of"):
        d.pattern("mosaic", "x", _photo(tmp_path))


def test_bleed_count_feeds_l24(tmp_path):
    d = Deck({**STYLE, "imagery": {"max_bleed": 1}})
    d.pattern("bleed-panel", "a", _photo(tmp_path))
    d.pattern("bleed-panel", "b", _photo(tmp_path))
    assert any("L24" in w for w in d.warnings)


# ---- PowerPoint render: spec section 9 ----
@pytest.mark.skipif(render.detect_backend() != "powerpoint", reason="needs PowerPoint")
def test_rendered_bleed_scrim_title_contrast(tmp_path):
    d = Deck(STYLE)
    d.pattern("bleed-scrim", "A quiet sentence over a photo", _two_tone(tmp_path), side="left")
    src = d.save(tmp_path / "t.pptx")
    im = Image.open(render.render(src, tmp_path / "r", width=1600)["slides"][0]).convert("RGB")
    x, w = d.col(0, 7)
    _, ty, _, th = d.style["canvas"]["title_box"]
    k = im.width / d.W
    region = im.crop((round(x * k), round(ty * k), round((x + w) * k), round((ty + th) * k)))
    text = L.hex_rgb(STYLE["color"][d.text_color])
    ratios = sorted(r for r in (L.contrast(text, p) for p in getattr(region, "get_flattened_data", region.getdata)()) if r >= 1.5)  # drop glyph pixels
    assert len(ratios) > 1000 and ratios[len(ratios) // 20] >= 3.0  # 5th percentile of the background


# ---- suggest() ----
def _names_of(res):
    return [r["pattern"] for r in res]


PHOTO = {"type": "photo", "size": [3000, 2000], "focus": [0.4, 0.3, 0.6, 0.7]}


def test_suggest_no_image_is_type_only():
    assert _names_of(suggest(None, "statement")) == ["type-only"]
    assert _names_of(suggest(PHOTO, "statement", n_images=0)) == ["type-only"]


def test_suggest_diagram_figure_screenshot_annotated_first():
    assert suggest({"type": "diagram", "size": [800, 600]}, "evidence")[0]["pattern"] == "figure"
    assert suggest({"type": "logo"}, "cover")[0]["pattern"] == "figure"
    shot = suggest({"type": "screenshot", "size": [1600, 900]}, "evidence")
    assert _names_of(shot) == ["annotated", "inset"]


def test_suggest_three_images_gallery():
    assert suggest(PHOTO, "comparison", n_images=3)[0]["pattern"] == "gallery"


def test_suggest_panorama_strip():
    res = suggest({"type": "photo", "size": [4000, 1000]}, "evidence")
    assert res[0]["pattern"] == "strip" and "4.0:1" in res[0]["reason"]


def test_suggest_cover_photo_scrim_then_panel_and_wordy_panel_first():
    assert _names_of(suggest(PHOTO, "cover", words=5)) == ["bleed-scrim", "bleed-panel"]
    assert _names_of(suggest(PHOTO, "cover", words=20))[0] == "bleed-panel"


def test_suggest_bleed_exhausted_puts_split_first():
    res = suggest(PHOTO, "cover", bleed_used=3)
    assert res[0]["pattern"] == "split"
    assert suggest(PHOTO, "cover", bleed_used=2)[0]["pattern"] == "bleed-scrim"
    assert suggest(PHOTO, "cover", bleed_used=1, imagery={"patterns": ["bleed-scrim", "split"], "max_bleed": 1})[0]["pattern"] == "split"


def test_suggest_skips_pattern_used_on_both_previous_slides():
    assert "split" not in _names_of(suggest(PHOTO, "evidence", prev=["split", "split"]))
    assert "split" in _names_of(suggest(PHOTO, "evidence", prev=["inset", "split"]))
    assert "split" in _names_of(suggest(PHOTO, "evidence", prev=["split"]))


def test_suggest_filters_disallowed_patterns_type_only_always_allowed():
    res = suggest(PHOTO, "cover", imagery={"patterns": ["inset"], "max_bleed": 3})
    assert _names_of(res) == ["inset", "type-only"]
    assert _names_of(suggest(PHOTO, "cover", imagery={"patterns": ["gallery"], "max_bleed": 3})) == ["type-only"]


def test_suggest_focus_left_text_right():
    left = suggest({**PHOTO, "focus": [0.0, 0.0, 0.3, 1.0]}, "cover")
    assert left[0]["params"]["side"] == "right"
    assert suggest({**PHOTO, "focus": [0.7, 0.0, 1.0, 1.0]}, "cover")[0]["params"]["side"] == "left"
    assert suggest({"type": "photo", "size": [3000, 2000]}, "cover")[0]["params"]["side"] == "left"


def test_suggest_tall_photo_split_ratio():
    res = suggest({"type": "photo", "size": [1000, 1600]}, "evidence")
    assert res[0]["pattern"] == "split" and res[0]["params"]["ratio"] == (4, 8)


@pytest.mark.parametrize("kind", ["diagram", "logo"])
def test_suggest_never_offers_cover_crop_patterns_for_diagram_or_logo(kind):
    for role in ("cover", "evidence", "comparison", "detail"):
        got = _names_of(suggest({"type": kind, "size": [4000, 1000]}, role, n_images=3))
        assert got[0] == "figure" and not set(got) & {"split", "gallery", "strip", "bleed-scrim", "bleed-panel"}


def test_suggest_unknown_role_raises():
    with pytest.raises(ValueError, match="role must be one of"):
        suggest(PHOTO, "intro")


def test_suggest_results_run_through_the_helpers(d, tmp_path):
    img = _photo(tmp_path)
    entry = {"type": "photo", "size": [2400, 1350], "focus": list(FOCUS), "must_keep": list(KEEP)}
    built = 0
    for role in ("cover", "section", "statement", "evidence", "comparison", "detail"):
        for words in (3, 20):
            for sug in suggest(entry, role, words=words):
                if sug["pattern"] in ("figure", "annotated", "type-only"):
                    continue
                kw = {"caption": "c"} if sug["pattern"] == "inset" else {}
                target = [img, img] if sug["pattern"] == "gallery" else img
                before = len(d.prs.slides)
                d.pattern(sug["pattern"], "T", target, focus=entry["focus"], must_keep=entry["must_keep"],
                          words=words, **kw, **sug["params"])
                assert len(d.prs.slides) == before + 1
                built += 1
    assert built >= 12


# ---- role kwarg: a cover can use the head size ----
HEAD = STYLE.get("scale", {}).get("head", 44)


def _title_size_and_box(slide):
    t = slide.shapes.title
    return t.text_frame.paragraphs[0].runs[0].font.size.pt, t.height / EMU


@pytest.mark.parametrize("name,kw", [("bleed-scrim", {"side": "left"}), ("bleed-panel", {"side": "left"}),
                                     ("split", {"side": "left"}), ("split", {"side": "left", "fit": "contain"})])
def test_role_head_gives_a_title_box_tall_enough_for_two_lines(d, tmp_path, name, kw):
    size = d.style["scale"]["head"]
    s, box = d.pattern(name, "승인 단계\n줄이기", _photo(tmp_path), role="head", **kw)
    pt, h = _title_size_and_box(s)
    assert pt == size and h >= 2 * 1.2 * size * d.style["rhythm"]["line_height"] / 72
    ty = d.style["canvas"]["title_box"][1]
    if box is not None:
        assert box[1] >= ty + h  # body text starts below the taller title


def test_default_role_keeps_the_title_box(d, tmp_path):
    s, _ = d.pattern("bleed-scrim", "x", _photo(tmp_path))
    pt, h = _title_size_and_box(s)
    assert pt == d.style["scale"]["title"] and h == pytest.approx(d.style["canvas"]["title_box"][3])


def test_bleed_scrim_plans_legibility_at_the_role_size(d, tmp_path, monkeypatch):
    seen = []
    real = L.plan
    monkeypatch.setattr(L, "plan", lambda px, ink, bg, size: seen.append(size) or real(px, ink, bg, size))
    d.pattern("bleed-scrim", "x", _photo(tmp_path), role="head")
    d.pattern("bleed-scrim", "x", _photo(tmp_path), role="head", words=5)
    assert seen == [d.style["scale"]["head"], d.style["scale"]["body"]]


def test_unknown_role_raises(d, tmp_path):
    with pytest.raises(ValueError, match="role"):
        d.pattern("bleed-panel", "x", _photo(tmp_path), role="huge")


# ---- bleed patterns keep clear of the image focus ----
from deckkit import patterns as P  # noqa: E402

LEFT_FOCUS = (0.0, 0.2, 0.3, 0.8)


def _wrong_default_side(monkeypatch):
    """Force the default text side onto the focus, to exercise the other-side retry."""
    monkeypatch.setattr(P, "_side", lambda side, focus: side or "left")


def test_focus_cover_maps_focus_through_the_actual_crop(d, tmp_path):
    img = _photo(tmp_path, size=(3000, 1000))  # 16:9 window is 1778px wide, centred
    box = (0, 0, d.W / 2, d.H)
    full = (0.0, 0.0, 1.0, 1.0)
    assert P._focus_cover(d, img, None, None, box) == 0.0
    assert P._focus_cover(d, img, full, None, box) == pytest.approx(0.5, abs=0.01)
    edge = (0.0, 0.0, 0.1, 1.0)  # the crop window starts at the focus: it sits in the left 2.2in
    assert P._focus_cover(d, img, edge, None, box) == 1.0
    assert P._focus_cover(d, img, edge, None, (d.W / 2, 0, d.W / 2, d.H)) == 0.0
    assert P._focus_cover(d, img, (0.45, 0.2, 0.55, 0.8), None, box) == pytest.approx(0.5, abs=0.01)


def test_bleed_panel_tries_the_other_side_when_the_panel_covers_the_focus(d, tmp_path, monkeypatch):
    _wrong_default_side(monkeypatch)
    s, box = d.pattern("bleed-panel", "T", _photo(tmp_path), focus=LEFT_FOCUS)
    panel = next(sh for sh in s.shapes if sh.name == "pc:panel")
    assert panel.left > 0 and box[0] > d.W / 2  # text and panel moved to the right
    assert not any("covers the image focus" in w for w in d.warnings)


def test_bleed_panel_explicit_side_is_kept_and_warns(d, tmp_path):
    s, box = d.pattern("bleed-panel", "T", _photo(tmp_path), focus=LEFT_FOCUS, side="left")
    panel = next(sh for sh in s.shapes if sh.name == "pc:panel")
    assert panel.left == 0 and (box[0], box[2]) == pytest.approx(d.col(0, 5))
    assert any("p.jpg: the panel covers the image focus; use split or pass a focus that sits clear of the text side" in w
               for w in d.warnings)


def test_bleed_panel_uses_fewer_columns_when_both_sides_cover_the_focus(d, tmp_path):
    s, box = d.pattern("bleed-panel", "T", _photo(tmp_path), focus=(0.4, 0.2, 0.6, 0.8))
    assert (box[0], box[2]) == pytest.approx(d.col(0, 4))
    assert next(sh for sh in s.shapes if sh.name == "pc:panel").left == 0
    assert not any("covers the image focus" in w for w in d.warnings)


def test_bleed_panel_without_focus_is_unchanged(d, tmp_path):
    s, box = d.pattern("bleed-panel", "T", _photo(tmp_path))
    assert (box[0], box[2]) == pytest.approx(d.col(0, 5)) and not d.warnings


def test_bleed_scrim_tries_the_other_side_when_the_band_covers_the_focus(d, tmp_path, monkeypatch):
    _wrong_default_side(monkeypatch)
    s, _ = d.pattern("bleed-scrim", "T", _two_tone(tmp_path), focus=LEFT_FOCUS)
    scrim = next(sh for sh in s.shapes if sh.name == "pc:scrim")
    assert scrim.left > 0 and not any("covers the image focus" in w for w in d.warnings)


def test_bleed_scrim_explicit_side_is_kept_and_warns(d, tmp_path):
    s, _ = d.pattern("bleed-scrim", "T", _two_tone(tmp_path), focus=LEFT_FOCUS, side="left")
    scrim = next(sh for sh in s.shapes if sh.name == "pc:scrim")
    assert scrim.left == 0
    assert any("n.png: the text band covers the image focus" in w for w in d.warnings)


def test_bleed_scrim_without_focus_is_unchanged(d, tmp_path):
    s, _ = d.pattern("bleed-scrim", "T", _two_tone(tmp_path))
    assert not any("focus" in w for w in d.warnings)
