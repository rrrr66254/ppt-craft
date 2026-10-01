"""deckkit: a thin guard against easy mistakes, on top of python-pptx.

build.py designs the layout every time. deckkit only handles what is easy to get wrong:
removing theme effects (shadows, gradients), Hangul fonts (ea), word-level line breaking (ko-KR), tracking,
automatic text-size fitting, grid coordinates, native cropping, native charts.
"""
import re
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from .crop import contain_box, cover_crop, effective_dpi
from .fonts import find_font_file
from .style import load_style
from .svg import svg_aspect, svg_is_safe
from .textfit import fit_size, has_broken_word

__all__ = ["Deck", "load_style"]

NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
SVG_EXT_URI = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"
NS_ASVG = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
HANGUL = re.compile(r"[\u3130-\u318f\uac00-\ud7a3]")
MIN_PT = {"head": 28, "title": 20, "governing": 16, "body": 14, "caption": 10, "mono": 12}
HEAD_ROLES = {"head", "title"}
BOLD_ROLES = {"head", "title"}
_ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
_ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
_TITLE_ONLY = 5  # the "Title Only" layout of the python-pptx default template
CHART_TYPES = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED,
               "line": XL_CHART_TYPE.LINE_MARKERS}
SERIES_COLORS = ("accent", "ink", "muted")
_SOLID = f'<a:solidFill xmlns:a="{NS_A}"><a:schemeClr val="phClr"/></a:solidFill>'
_NO_EFFECT = f'<a:effectStyle xmlns:a="{NS_A}"><a:effectLst/></a:effectStyle>'


class Deck:
    def __init__(self, style, lang="ko"):
        self.style = load_style(style)
        self.lang = lang
        self.W, self.H = self.style["canvas"]["size"]
        self.m = self.style["canvas"]["margin"]
        self.warnings = []
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(self.W), Inches(self.H)
        self._neutralize_theme(self.prs.slide_master.part.part_related_by(RT.THEME))
        self._notes_themed = False
        for role in ("head", "body", "mono"):
            family = self.style["font"][role]
            if find_font_file(family) is None:
                self._warn(f"Font '{family}' is not installed. Rendering will substitute another font.")

    # ---- coordinates ----
    def col(self, start, span, cols=12, gutter=0.25):
        """12-column grid: (x, w) in inches for span columns starting at column `start` (0-based)."""
        inner = self.W - 2 * self.m
        cw = (inner - gutter * (cols - 1)) / cols
        return (self.m + start * (cw + gutter), span * cw + (span - 1) * gutter)

    @property
    def content_top(self):
        _, y, _, h = self.style["canvas"]["title_box"]
        return y + h + 0.3

    @property
    def content_bottom(self):
        return self.H - self.m - 0.35  # room for the footer

    def body_box(self, start=0, span=12, top=None, bottom=None):
        x, w = self.col(start, span)
        t = self.content_top if top is None else top
        b = self.content_bottom if bottom is None else bottom
        return (x, t, w, b - t)

    # ---- slides ----
    def slide(self, title=None, *, bg="bg", box=None, role="title", color="ink", anchor="bottom"):
        """New slide. The title goes into the title placeholder (accessibility, outline view), pinned to title_box (or box)."""
        self._check(role=role, anchor=anchor)
        self._rgb(bg), self._rgb(color)
        s = self.prs.slides.add_slide(self.prs.slide_layouts[_TITLE_ONLY])
        fill = s.background.fill
        fill.solid()
        fill.fore_color.rgb = self._rgb(bg)
        ph = s.shapes.title
        if title is None:
            ph._element.getparent().remove(ph._element)
            return s
        box = box or self.style["canvas"]["title_box"]
        ph.left, ph.top, ph.width, ph.height = (Inches(v) for v in box)
        self._fill(ph, str(title).split("\n"), role, color=color, align="left", anchor=anchor, fit=True, box=box)
        return s

    def text(self, slide, box, content, role="body", *, color="ink", bold=None, align="left",
             anchor="top", fit=True, bullets=False, size=None):
        """Text box. content is a string or a list of paragraphs. bullets=True sets the paragraph bullet property (no literal • character).
        size: pt to use instead of the role default (big numbers, etc.). With fit=True the text may shrink to fit the box."""
        self._check(role=role, align=align, anchor=anchor)
        if size is not None and not (isinstance(size, (int, float)) and not isinstance(size, bool) and size > 0):
            raise ValueError(f"size must be a number greater than 0 (pt) (got: {size!r})")
        self._rgb(color)
        paragraphs = [content] if isinstance(content, str) else list(content)
        paragraphs = [line for p in paragraphs for line in str(p).split("\n")]
        shape = slide.shapes.add_textbox(*(Inches(v) for v in box))
        self._fill(shape, paragraphs, role, color=color, align=align, anchor=anchor, fit=fit, box=box,
                   bold=bold, bullets=bullets, size=size)
        return shape

    def rect(self, slide, box, color="muted"):
        """Solid-color rectangle. No border, no shadow."""
        rgb = self._rgb(color)
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, *(Inches(v) for v in box))
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb
        shape.line.fill.background()
        _no_effects(shape)
        return shape

    def rule(self, slide, x, y, w, *, color="ink", weight=0.75):
        """Horizontal line. Do not use it as an accent rule right under the title (L6)."""
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x + w), Inches(y))
        line.line.color.rgb = self._rgb(color)
        line.line.width = Pt(weight)
        _no_effects(line)
        return line

    def send_to_back(self, slide, shape):
        """Send a shape to the very back (so the title shows over a background block)."""
        tree = slide.shapes._spTree
        tree.remove(shape._element)
        tree.insert(2, shape._element)

    def chart(self, slide, box, kind, categories, series, *, highlight=None, number_format="General", source=None,
              labels=False):
        """Native chart (editable). series={name: list of values}. highlight=index of the category to emphasize (bars, single series).
        labels=True: put value labels on bars/points and hide the value axis."""
        if kind not in CHART_TYPES:
            raise ValueError(f"kind must be one of {list(CHART_TYPES)}")
        if len(series) > len(SERIES_COLORS):
            self._warn("With more than 3 series the colors are hard to tell apart. Split the chart.")
        cats = list(categories)
        if highlight is not None and not 0 <= highlight < len(cats):
            raise ValueError(f"highlight must be a category index from 0 to {len(cats) - 1} (got: {highlight})")
        if highlight is not None and len(series) > 1:
            self._warn("highlight applies only to the first series. Do not use highlight with multiple series.")
        data = CategoryChartData()
        data.categories = cats
        for name, values in series.items():
            data.add_series(name, list(values))
        frame = slide.shapes.add_chart(CHART_TYPES[kind], *(Inches(v) for v in box), data)
        chart = frame.chart
        chart.has_title = False  # The slide title carries the claim. No chart title.
        body = self.style["font"]["body"]
        chart.font.size = Pt(self.style["scale"]["caption"])
        chart.font.name = body
        chart.font.color.rgb = self._rgb("muted")
        rpr = chart._chartSpace.find(f".//{qn('c:txPr')}//{qn('a:defRPr')}")
        ea = rpr.find(qn("a:ea"))
        if ea is None:
            ea = etree.Element(qn("a:ea"))
            rpr.find(qn("a:latin")).addnext(ea)
        ea.set("typeface", body)
        if self.lang == "ko":
            rpr.set("lang", "ko-KR")
        chart.has_legend = len(series) > 1
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.TOP
            chart.legend.include_in_layout = False
        plot = chart.plots[0]
        va = chart.value_axis
        va.has_major_gridlines = False
        va.tick_labels.number_format = number_format
        va.tick_labels.number_format_is_linked = False
        va.format.line.fill.background()
        if kind != "line" and all(v >= 0 for vals in series.values() for v in vals if v is not None):
            va.minimum_scale = 0  # bars encode value by length: never let auto-scaling exaggerate small differences (not with negatives)
        if labels:
            plot.has_data_labels = True
            dl = plot.data_labels
            dl.show_value = True
            dl.number_format = number_format
            dl.number_format_is_linked = False
            dl.font.size = Pt(self.style["scale"]["caption"])
            dl.font.color.rgb = self._rgb("ink")
            dl.position = XL_LABEL_POSITION.ABOVE if kind == "line" else XL_LABEL_POSITION.OUTSIDE_END
            va.visible = False
            va._element.find(qn("c:delete")).set("val", "1")  # explicit instead of omitting the default
        chart.category_axis.format.line.color.rgb = self._rgb("muted")
        if kind != "line":
            plot.gap_width = 60
        if kind == "bar":  # flip so the first category is on top, keep the value-axis ticks at the bottom
            chart.category_axis.reverse_order = True
            chart.value_axis._element.find(qn("c:crosses")).set("val", "max")  # the value axis crosses at the last (= bottom) category
        for i, s in enumerate(plot.series):
            rgb = self._rgb(SERIES_COLORS[min(i, len(SERIES_COLORS) - 1)])
            if kind == "line":
                s.format.line.color.rgb = rgb
                s.format.line.width = Pt(2.25)
                s.marker.style = XL_MARKER_STYLE.CIRCLE
                s.marker.format.fill.solid()
                s.marker.format.fill.fore_color.rgb = rgb
                s.marker.format.line.color.rgb = rgb
            else:
                s.format.fill.solid()
                s.format.fill.fore_color.rgb = rgb
        if highlight is not None and kind != "line":
            for j in range(len(cats)):
                point = plot.series[0].points[j]
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = self._rgb("accent" if j == highlight else "muted")
        if source:
            label = "출처" if self.lang == "ko" else "Source"  # chart source label, in the deck's language
            self.text(slide, (box[0], box[1] + box[3] + 0.08, box[2], 0.3), f"{label}: {source}", "caption",
                      color="muted", fit=False)
        return frame

    def image(self, slide, box, path, *, fit="cover", focus=None, must_keep=None, anchor="center"):
        """Place an image with PowerPoint native cropping (original preserved, can be re-cropped in PowerPoint).

        fit="cover": crop to fill the slot (around focus). If must_keep cannot be honored, fall back to contain.
        fit="contain": no cropping, fit inside the slot (diagrams, logos).
        """
        path = Path(path)
        with Image.open(path) as im:
            iw, ih = im.size
            if im.getexif().get(0x0112, 1) != 1:
                self._warn(f"{path.name}: has EXIF rotation info. Normalize it with images.py first.")
        x, y, w, h = box
        crop = (0.0, 0.0, 0.0, 0.0)
        if fit == "cover":
            crop = cover_crop(iw, ih, w, h, focus, must_keep, anchor)
            if crop is None:
                self._warn(f"{path.name}: the must_keep region cannot be kept at the {w:.1f}x{h:.1f}in slot ratio; fell back to contain.")
                fit, crop = "contain", (0.0, 0.0, 0.0, 0.0)
        if fit == "contain":
            x, y, w, h = contain_box(iw, ih, box)
        pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
        pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = crop
        dpi = effective_dpi(iw, crop, w)
        if dpi < 150:
            self._warn(f"{path.name}: resolution for the slot is {dpi:.0f}dpi (under 150), so it may look blurry.")
        return pic

    def icon(self, slide, box, svg_path, png_path=None):
        """Place an icon as SVG (PowerPoint 2019+/365 shows the vector; older viewers show the PNG fallback).
        png_path defaults to svg_path with .png; it is rendered with resvg-py if missing. Keeps aspect (contain)."""
        svg_path = Path(svg_path)
        png_path = Path(png_path) if png_path else svg_path.with_suffix(".png")
        svg = svg_path.read_bytes()
        try:
            text = svg.decode("utf-8")
        except UnicodeDecodeError:
            raise ValueError(f"{svg_path.name}: the SVG is not UTF-8 text.") from None
        if not svg_is_safe(text):
            raise ValueError(f"{svg_path.name}: the SVG contains external references or scripts and was refused.")
        aspect = svg_aspect(text)
        if not png_path.exists():
            try:
                import resvg_py
            except ImportError:
                raise RuntimeError("resvg-py is missing: pip install -r requirements.txt") from None
            size = {} if aspect is None else {"width": 512} if aspect >= 1 else {"height": 512}  # longest side 512
            try:
                png_path.write_bytes(bytes(resvg_py.svg_to_bytes(svg_string=text, **size)))
            except Exception:  # noqa: BLE001 - resvg raises a plain error on unusable SVG
                raise ValueError(f"{svg_path.name}: could not render") from None
        with Image.open(png_path) as im:
            iw, ih = im.size
        if aspect is not None and abs(iw / ih - aspect) / aspect > 0.02:
            self._warn(f"{png_path.name}: PNG aspect ratio differs from {svg_path.name} by more than 2%, so the "
                       f"fallback will look different from the vector.")
        x, y, w, h = contain_box(iw, ih, box)
        pic = slide.shapes.add_picture(str(png_path), Inches(x), Inches(y), Inches(w), Inches(h))
        pkg = slide.part.package
        svg_part = Part(pkg.next_partname("/ppt/media/image%d.svg"), "image/svg+xml", pkg, svg)
        rid = slide.part.relate_to(svg_part, RT.IMAGE)
        blip = pic._element.find(f".//{qn('a:blip')}")
        ext_lst = blip.find(qn("a:extLst"))
        if ext_lst is None:
            ext_lst = etree.SubElement(blip, qn("a:extLst"))
        ext = etree.SubElement(ext_lst, qn("a:ext"))
        ext.set("uri", SVG_EXT_URI)
        etree.SubElement(ext, f"{{{NS_ASVG}}}svgBlip", nsmap={"asvg": NS_ASVG}).set(qn("r:embed"), rid)
        return pic

    def callout(self, slide, target_box, label, label_box, *, color="accent"):
        """Annotation on a screenshot (H3): border around the target area + connector line + short note. All native shapes."""
        target = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, *(Inches(v) for v in target_box))
        target.fill.background()
        target.line.color.rgb = self._rgb(color)
        target.line.width = Pt(2)
        _no_effects(target)
        lx, ly, lw, lh = label_box
        tx, ty, tw, th = target_box
        align = "left"
        # pick the side with the biggest gap between label and target (ties: right, left, above, below)
        gaps = {"right": lx - (tx + tw), "left": tx - (lx + lw), "above": ty - (ly + lh), "below": ly - (ty + th)}
        side = max(gaps, key=gaps.get)
        if side == "right":
            sx, sy, ex, ey = lx, ly + lh / 2, tx + tw, ty + th / 2
        elif side == "left":
            sx, sy, ex, ey = lx + lw, ly + lh / 2, tx, ty + th / 2
            align = "right"
        elif side == "above":   # label is above
            sx, sy, ex, ey = lx + lw / 2, ly + lh, tx + tw / 2, ty
        else:                   # label is below
            sx, sy, ex, ey = lx + lw / 2, ly, tx + tw / 2, ty + th
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(sx), Inches(sy), Inches(ex), Inches(ey))
        line.line.color.rgb = self._rgb(color)
        line.line.width = Pt(1.5)
        _no_effects(line)
        self.text(slide, label_box, label, "caption", color=color, align=align, anchor="middle")
        return target

    def notes(self, slide, text):
        slide.notes_slide.notes_text_frame.text = text
        if not self._notes_themed:  # the notes master's theme (theme2) also has shadows and gradients
            self._neutralize_theme(self.prs.notes_master.part.part_related_by(RT.THEME))
            self._notes_themed = True

    def footer(self, slide, left="", page=None):
        """Footer: event name, date, etc. on the left, page number on the right (H4)."""
        y = self.H - self.m - 0.3
        if left:
            self.text(slide, (self.m, y, self.W * 0.6, 0.3), left, "caption", color="muted", fit=False)
        if page is not None:
            self.text(slide, (self.W - self.m - 1.0, y, 1.0, 0.3), str(page), "caption", color="muted",
                      align="right", fit=False)

    def save(self, path):
        path = Path(path)
        self.prs.save(str(path))
        for w in self.warnings:
            line = f"[deckkit warning] {w}"
            try:
                print(line)
            except UnicodeEncodeError:  # narrow consoles such as cp949
                print(line.encode("ascii", "backslashreplace").decode())
        return path

    # ---- internals ----
    def _warn(self, msg):
        if msg not in self.warnings:
            self.warnings.append(msg)

    def _check(self, role=None, align=None, anchor=None):
        if role is not None and role not in MIN_PT:
            raise ValueError(f"role '{role}' must be one of {list(MIN_PT)}")
        if align is not None and align not in _ALIGN:
            raise ValueError(f"align '{align}' must be one of {list(_ALIGN)}")
        if anchor is not None and anchor not in _ANCHOR:
            raise ValueError(f"anchor '{anchor}' must be one of {list(_ANCHOR)}")

    def _rgb(self, name):
        value = self.style["color"].get(name, name)
        if not re.fullmatch(r"#[0-9A-Fa-f]{6}", str(value)):
            raise ValueError(f"color '{name}' must be a style color name (bg/ink/muted/accent) or #RRGGBB")
        return RGBColor.from_string(value[1:].upper())

    def _fill(self, shape, paragraphs, role, *, color, align, anchor, fit, box, bold=None, bullets=False,
              size=None):
        st = self.style
        family = st["font"]["mono" if role == "mono" else "head" if role in HEAD_ROLES else "body"]
        size = st["scale"][role] if size is None else size
        spacing = st["rhythm"]["line_height"]
        tracking = st["rhythm"]["tracking_head" if role in HEAD_ROLES else "tracking_body"]
        is_bold = (role in BOLD_ROLES) if bold is None else bold
        indent = size * 1.1 if bullets else 0.0
        multi = len(paragraphs) > 1
        gap = 0.0 if role in HEAD_ROLES else 0.5  # no paragraph spacing between title lines
        if fit:
            ref = find_font_file(family, is_bold)
            size, ok = fit_size(paragraphs, ref, size, MIN_PT[role], box[2], box[3], spacing, tracking,
                                para_gap_ratio=gap if multi else 0.0, indent_pt=indent)
            if not ok:
                self._warn(f"Text overflow: '{paragraphs[0][:24]}' - still overflows the {box[2]:.2f}x{box[3]:.2f}in "
                           f"box at {size}pt. Cut the content or enlarge the box.")
            if has_broken_word(paragraphs, ref, size, box[2] - indent / 72, tracking):
                self._warn(f"Word broken mid-line: '{paragraphs[0][:24]}' - a word is longer than the box.")
        tf = shape.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = _ANCHOR[anchor]
        tf.clear()
        for i, text in enumerate(paragraphs):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = _ALIGN[align]
            p.line_spacing = spacing
            if multi and gap and i < len(paragraphs) - 1:
                p.space_after = Pt(size * gap)
            if bullets:
                _bullet(p, indent)
            run = p.add_run()
            run.text = text
            self._run_props(run, family, size, is_bold, color, tracking)

    def _run_props(self, run, family, size, bold, color, tracking):
        font = run.font
        font.size = Pt(size)
        font.bold = bold
        font.color.rgb = self._rgb(color)
        font.name = family
        rPr = run._r.get_or_add_rPr()
        ea = rPr.find(qn("a:ea"))
        if ea is None:
            ea = etree.Element(qn("a:ea"))
            rPr.find(qn("a:latin")).addnext(ea)
        ea.set("typeface", family)
        rPr.set("spc", str(round(tracking * size * 100)))
        if HANGUL.search(run.text):
            rPr.set("lang", "ko-KR")
            rPr.set("altLang", "en-US")

    def _neutralize_theme(self, part):
        """Remove gradient fills, shadows, and 3D effects from the theme part, and set color and font to the style values."""
        in_place = hasattr(part, "_element")  # the notes-master theme is an XmlPart, the slide-master theme is a Part (blob)
        root = part._element if in_place else etree.fromstring(part.blob)

        def a(tag):
            return f"{{{NS_A}}}{tag}"

        fmt = root.find(f".//{a('fmtScheme')}")
        for tag in ("fillStyleLst", "bgFillStyleLst"):
            fmt.find(a(tag))[:] = [etree.fromstring(_SOLID) for _ in range(3)]
        fmt.find(a("effectStyleLst"))[:] = [etree.fromstring(_NO_EFFECT) for _ in range(3)]
        font = self.style["font"]
        for tag, family in (("majorFont", font["head"]), ("minorFont", font["body"])):
            node = root.find(f".//{a(tag)}")
            node.find(a("latin")).set("typeface", family)
            node.find(a("ea")).set("typeface", family)
            for script in node.findall(a("font")):
                if script.get("script") == "Hang":
                    script.set("typeface", family)
        color = self.style["color"]
        scheme = root.find(f".//{a('clrScheme')}")
        for tag, key in (("dk1", "ink"), ("lt1", "bg"), ("dk2", "ink"), ("lt2", "bg"),
                         ("accent1", "accent"), ("accent2", "muted"), ("accent3", "muted"),
                         ("accent4", "muted"), ("accent5", "muted"), ("accent6", "muted")):
            scheme.find(a(tag))[:] = [etree.fromstring(f'<a:srgbClr xmlns:a="{NS_A}" val="{color[key][1:].upper()}"/>')]
        if not in_place:
            part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _no_effects(shape):
    """Cut the theme effect reference (shadows, etc.): put an explicit empty effectLst in spPr."""
    spPr = shape._element.spPr
    if spPr.find(qn("a:effectLst")) is not None:
        return
    eff = etree.SubElement(spPr, qn("a:effectLst"))
    for tag in ("a:scene3d", "a:sp3d"):  # schema order: ln, effectLst, scene3d, sp3d
        nxt = spPr.find(qn(tag))
        if nxt is not None:
            nxt.addprevious(eff)
            break


def _bullet(p, indent_pt):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Pt(indent_pt))))
    pPr.set("indent", str(-int(Pt(indent_pt))))
    for tag in ("a:buNone", "a:buAutoNum", "a:buChar", "a:buFont"):
        for old in pPr.findall(qn(tag)):
            pPr.remove(old)
    etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
    etree.SubElement(pPr, qn("a:buChar")).set("char", "•")
