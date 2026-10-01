"""Automated pre-render check: finds the AI-tell rules that text and XML alone can catch (rule IDs are in rules/tells.md).
Usage: python lint.py <pptx> [--style style.json] [--out lint.json]
Exit codes: 0 clean (0 even with only minors), 1 blockers/majors present, 2 lint failed (corrupt file, etc.).
Items that can only be judged from captures are left to deck-reviewer."""
import argparse
import json
import re
import statistics
import sys
from collections import Counter
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
WORDS = json.loads((ROOT / "rules" / "banned-words.json").read_text(encoding="utf-8"))

SEVERITY = {"W14": "blocker", "I3": "blocker", "T12": "blocker",
            "L2": "major", "L6": "major", "C1": "major", "C6": "major", "D3": "major",
            "D5": "major", "S1": "major", "T11": "major"}
ORDER = {"blocker": 0, "major": 1, "minor": 2}
FONT_RULE = {"inter": "T1", "roboto": "T1", "arial": "T1",
             "poppins": "T2", "montserrat": "T2", "space grotesk": "T2", "instrument serif": "T2", "geist": "T2",
             "aptos": "T8", "calibri": "T8"}
OFFICE_KO = {"맑은 고딕", "malgun gothic"}  # the default theme's Hang font: does not count as specified
ARROWS = re.compile("[→⇒➔➜➡≈✓✔]")
EMOJI = re.compile("[\U0001F1E6-\U0001F1FF\U0001F300-\U0001FAFF]"
                   "|[\u2705\u274C\u274E\u2753-\u2755\u2757\u26A1\u2728\u23F0-\u23F3\u2B50\u2B55]"
                   "|[\u2600-\u27BF]\uFE0F|\u20E3")
HANGUL = re.compile(r"[\uac00-\ud7a3]")
NUMBER_CLAIM = re.compile(r"\d+(?:[.,]\d+)?\s?%|\d+(?:\.\d+)?\s?[xX]\b|\d+(?:\.\d+)?배(?!포)"
                          r"|\d+\s?(?:억|만)\s?(?:원|명)?|[$₩€]\s?\d")
SOURCE_MARK = re.compile(r"출처|\bsources?\s*[:：]|\bsource\b|자료\s?:|참고\s?:|\[출처 필요\]", re.I)
PLACEHOLDER = re.compile(r"\[(?:insert|add|your|placeholder)[^\]]*\]|lorem ipsum|add your [^\n.]{0,30}\bhere\b"
                         r"|여기에 (?:입력|내용|텍스트)|텍스트를 입력하(?:십시오|세요)|\bTBD\b"
                         r"|(?-i:(?<![\w-])XXX(?![\w-]))", re.I)
TEMPLATE_TITLE = re.compile(r"^(agenda|목차|contents|key takeaways?|takeaways|thank you!?|thanks!?"
                            r"|감사합니다\.?|q\s?&\s?a|questions\??|질의\s?응답)$", re.I)
# Credit lines from `assets.py credits` look like "keyword: explanation" but are not prose (W8 does not apply to them).
CREDIT_LINE = re.compile(r"^(Photo|Icons?|AI-generated|사진|AI 생성)(?![A-Za-z])")
LEAD_IN = re.compile(r"^(?!\d)(?!(?:출처|source|자료|참고|주|note)\s*[:：])[^:：]{1,24}[:：](?!//)\s*\S", re.I)
TIME_COLON = re.compile(r"\d:\d")
SMALL_WORDS = {"a", "an", "the", "and", "or", "of", "to", "in", "on", "for", "with", "vs", "at", "by"}
BULLET_CHARS = ("•", "●", "▪", "■", "◦")
EMU = 914400
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _in(v):
    return (v or 0) / EMU


SHAPE_TAGS = tuple(qn(t) for t in ("p:sp", "p:pic", "p:graphicFrame", "p:grpSp", "p:cxnSp"))


def _top_shapes(slide):
    """Top-level shapes. Skips elements python-pptx cannot read (contentPart, etc.)."""
    for el in slide.shapes._spTree.iterchildren():
        if el.tag in SHAPE_TAGS:
            try:
                yield slide.shapes._shape_factory(el)
            except Exception:
                continue


IDENTITY = (1.0, 0.0, 1.0, 0.0)  # x' = ax*x + bx, y' = ay*y + by


def _geo(sh, xf):
    ax, bx, ay, by = xf
    return {"x": ax * _in(sh.left) + bx, "y": ay * _in(sh.top) + by,
            "w": ax * _in(sh.width), "h": ay * _in(sh.height)}


def _child_xf(el, g, parent_xf):
    """Transform that maps group-child coordinates (relative to chOff/chExt) into slide coordinates. g is the group box in slide coordinates."""
    xfrm = el.find(f"{qn('p:grpSpPr')}/{qn('a:xfrm')}")
    ch_off = xfrm.find(qn("a:chOff")) if xfrm is not None else None
    ch_ext = xfrm.find(qn("a:chExt")) if xfrm is not None else None
    if ch_off is None or ch_ext is None:
        return parent_xf
    cx, cy = int(ch_ext.get("cx", 0)) / EMU, int(ch_ext.get("cy", 0)) / EMU
    ax = g["w"] / cx if cx else 1.0  # scale 1 if chExt is 0
    ay = g["h"] / cy if cy else 1.0
    return (ax, g["x"] - int(ch_off.get("x", 0)) / EMU * ax, ay, g["y"] - int(ch_off.get("y", 0)) / EMU * ay)


def _walk(sh, xf=IDENTITY):
    """Yield (shape, slide-coordinate box) in turn, down to group descendants. Unreadable shapes are skipped."""
    try:
        geo = _geo(sh, xf)
    except Exception:
        return
    yield sh, geo
    if sh._element.tag == qn("p:grpSp"):
        try:
            cxf, children = _child_xf(sh._element, geo, xf), list(sh.shapes)
        except Exception:
            return
        for child in children:
            yield from _walk(child, cxf)


def _kind(sh):
    el = sh._element
    if el.tag == qn("p:grpSp"):
        return "group"
    if el.tag == qn("p:pic"):
        return "pic"
    if getattr(sh, "has_chart", False) and sh.has_chart:
        return "chart"
    if getattr(sh, "has_table", False) and sh.has_table:
        return "table"
    if el.tag == qn("p:cxnSp"):
        return "line"
    has_text = sh.has_text_frame and sh.text_frame.text.strip()
    c_nv = el.find(f"{qn('p:nvSpPr')}/{qn('p:cNvSpPr')}")
    if c_nv is not None and c_nv.get("txBox") == "1":
        return "text"
    geom = el.find(f"{qn('p:spPr')}/{qn('a:prstGeom')}")
    if geom is not None and geom.get("prst") == "line":
        return "line"
    if geom is not None and geom.get("prst") in ("rect", "roundRect"):
        return "text" if has_text else "rect"
    return "text" if has_text else "other"


def _is_title_ph(sh):
    try:
        return sh.is_placeholder and sh.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
    except Exception:
        return False


def _guess_title(tops):
    """When there is no title placeholder: the shape with the biggest text in the upper area (within 1.8in)."""
    best, best_size = None, 0
    for sh in tops:
        try:
            if not sh.has_text_frame or not sh.text_frame.text.strip() or _in(sh.top) > 1.8:
                continue
            size = max((r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size), default=0)
        except Exception:
            continue
        if size > best_size:
            best, best_size = sh, size
    return best


def _typeface(node, tag):
    child = node.find(f"{A}{tag}") if node is not None else None
    return child.get("typeface", "") if child is not None else ""


def _ko_font_set(theme, master):
    """Is a Hangul font specified in the deck: theme ea, a (non-default) Hang font, or ea in the master text styles."""
    for tag in ("minorFont", "majorFont"):
        node = theme.find(f".//{A}{tag}")
        if node is None:
            continue
        if _typeface(node, "ea"):
            return True
        for f in node.findall(f"{A}font"):
            face = f.get("typeface", "")
            if f.get("script") == "Hang" and face and face.lower() not in OFFICE_KO:
                return True
    for d in master.iterfind(f".//{qn('p:txStyles')}//{qn('a:defRPr')}"):
        face = _typeface(d, "ea")
        if face and not face.startswith("+"):
            return True
    return False


def _collect(tf, frame, is_title, bullet_default, paras, runs):
    for p in tf.paragraphs:
        for r in p.runs:
            rPr = r._r.find(qn("a:rPr"))
            runs.append({"text": r.text, "lang": rPr.get("lang", "") if rPr is not None else "",
                         "ea": _typeface(rPr, "ea"), "latin": _typeface(rPr, "latin")})
        if is_title or not p.text.strip():
            continue
        pPr = p._p.find(qn("a:pPr"))
        explicit = pPr is not None and pPr.find(qn("a:buChar")) is not None
        suppressed = pPr is not None and pPr.find(qn("a:buNone")) is not None
        paras.append({"text": p.text.strip(), "frame": frame,
                      "bold_runs": sum(1 for r in p.runs if r.font.bold),
                      "bullet": explicit or (bullet_default and not suppressed)})


def _inherits_bullets(sh):
    """Body and object placeholders inherit bullets from the layout."""
    if not sh.is_placeholder or sh.placeholder_format.type not in (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT):
        return False
    try:  # if the layout placeholder turns off level-1 bullets (section headers, etc.), it is not a bullet
        base = sh._base_placeholder
        lvl1 = base._element.find(f".//{qn('a:lstStyle')}/{qn('a:lvl1pPr')}") if base is not None else None
        if lvl1 is not None and lvl1.find(qn("a:buNone")) is not None:
            return False
    except Exception:
        pass
    return True


def load(path):
    prs = Presentation(str(path))
    theme = etree.fromstring(prs.slide_master.part.part_related_by(RT.THEME).blob)
    minor, major = theme.find(f".//{A}minorFont"), theme.find(f".//{A}majorFont")
    effects = theme.find(f".//{A}effectStyleLst")
    deck = {
        "ko_font_set": _ko_font_set(theme, prs.slide_master._element),
        "theme_fonts": {_typeface(major, "latin"), _typeface(minor, "latin")},
        # whether each effectStyle (numbered from 1) has a shadow or glow
        "effect_flags": [es.find(f".//{A}outerShdw") is not None or es.find(f".//{A}glow") is not None
                         for es in (effects.findall(f"{A}effectStyle") if effects is not None else [])],
        "height": _in(prs.slide_height),
        "slides": [],
    }
    for idx, s in enumerate(prs.slides, 1):
        tops = list(_top_shapes(s))
        title_sh = next((sh for sh in tops if _is_title_ph(sh)), None)
        if title_sh is None:
            title_sh = _guess_title(tops)
        title_id = title_sh.shape_id if title_sh is not None else None
        shapes, els, paras, runs = [], [], [], []
        frame = 0
        for top in tops:
            for sh, geo in _walk(top):  # group children are collected in slide coordinates, after the group transform
                try:
                    is_title = sh.shape_id == title_id
                    shapes.append({"kind": _kind(sh), **geo, "is_title": is_title, "el": sh._element})
                    els.append(sh._element)
                    if sh.has_text_frame:
                        frame += 1
                        _collect(sh.text_frame, frame, is_title, _inherits_bullets(sh), paras, runs)
                    elif getattr(sh, "has_table", False) and sh.has_table:
                        for row in sh.table.rows:
                            for cell in row.cells:
                                frame += 1
                                _collect(cell.text_frame, frame, False, False, paras, runs)
                except Exception:
                    continue
        title = title_sh.text_frame.text.strip() if title_sh is not None and title_sh.has_text_frame else ""
        notes = ""
        if s.has_notes_slide:
            tf = s.notes_slide.notes_text_frame  # None if there is no notes placeholder
            notes = tf.text if tf is not None else ""
        deck["slides"].append({"idx": idx, "title": title, "paras": paras, "runs": runs, "shapes": shapes,
                               "els": els, "xml": etree.tostring(s._element, encoding="unicode"), "notes": notes})
    return deck


def _f(out, rule, slide, msg):
    out.append({"rule": rule, "severity": SEVERITY.get(rule, "minor"), "slide": slide, "message": msg})


def _visible(sl):
    return "\n".join([sl["title"]] + [p["text"] for p in sl["paras"]])


def check_text(deck, out):
    per_deck = {rule: [] for rule in WORDS["per_deck"]}
    for sl in deck["slides"]:
        i, text = sl["idx"], _visible(sl)
        if PLACEHOLDER.search(text + "\n" + sl["notes"]):
            _f(out, "W14", i, "Leftover placeholder/instruction text")
        if EMOJI.search(ARROWS.sub("", text)):
            _f(out, "I3", i, "Emoji used")
        if len(ARROWS.findall(text)) >= 2:
            _f(out, "I9", i, "Chain of Unicode arrows/symbols")
        if "—" in text:
            _f(out, "W7", i, "em dash (—) used")
        for rule, patterns in WORDS["per_hit"].items():
            for pat in patterns:
                m = re.search(pat, text, re.I)
                if m:
                    _f(out, rule, i, f"Banned phrase '{m.group(0)}'")
                    break
        for rule, spec in WORDS["per_deck"].items():
            per_deck[rule] += [i for pat in spec["patterns"] for _ in re.finditer(pat, text)]
        if NUMBER_CLAIM.search(text) and not SOURCE_MARK.search(text + "\n" + sl["notes"]):
            _f(out, "D3", i, "Figure without a source: add a source on the slide or in the notes, or mark it [source needed]")
    for rule, hits in per_deck.items():
        if len(hits) >= WORDS["per_deck"][rule]["min"]:
            _f(out, rule, None, f"{len(hits)} times across the deck (slides {sorted(set(hits))})")


def check_titles(deck, out):
    for sl in deck["slides"]:
        t = sl["title"].replace("\n", " ").strip()
        if not t:
            continue
        if TEMPLATE_TITLE.match(t):
            _f(out, "S1", sl["idx"], f"Template-flow title '{t}'")
        if (":" in t or "：" in t) and not TIME_COLON.search(t):
            _f(out, "W3", sl["idx"], f"Colon title '{t}'")
        words = re.findall(r"[A-Za-z][A-Za-z'-]*", t)
        if not HANGUL.search(t) and len(words) >= 3 and all(w[0].isupper() for w in words if w.lower() not in SMALL_WORDS):
            _f(out, "T5", sl["idx"], f"Title Case title '{t}'")


def _signature(sl, height):
    """Layout fingerprint: leaves out the title and the footer (low shapes within the bottom 1in band)."""
    return tuple(sorted((s["kind"], round(s["x"]), round(s["y"]), round(s["w"]), round(s["h"]))
                        for s in sl["shapes"]
                        if not s["is_title"] and not (s["h"] <= 0.5 and s["y"] + s["h"] >= height - 1.0)))


def check_structure(deck, out):
    bullet_counts = []
    for sl in deck["slides"]:
        i, paras = sl["idx"], sl["paras"]
        bullets = [p for p in paras if p["bullet"]]
        if len(bullets) >= 2:
            bullet_counts.append(len(bullets))
        frames = {}
        for p in paras:
            frames.setdefault(p["frame"], []).append(len(p["text"]))
        for lens in frames.values():  # only look at 3 or more paragraphs inside a single text box
            if len(lens) >= 3 and statistics.mean(lens) >= 8 and statistics.pstdev(lens) / statistics.mean(lens) < 0.2:
                _f(out, "W9", i, "Item lengths are too uniform")
                break
        if sum(1 for p in paras if LEAD_IN.match(p["text"]) and not CREDIT_LINE.match(p["text"])) >= 2:
            _f(out, "W8", i, "Repeated 'keyword: explanation' pattern")
        if len(bullets) >= 3 and all(re.search(r"니다[.!]?$", p["text"]) for p in bullets):
            _f(out, "K15", i, "All bullets are formal '~니다' (-nida) declarative sentences")
        if sum(p["bold_runs"] for p in paras) >= 2:
            _f(out, "T9", i, "Bold overused in body text")
        if any(p["text"].startswith(BULLET_CHARS) for p in paras):
            _f(out, "BULLET", i, "Bullet character typed directly as text (risk of double bullets)")
    if len(bullet_counts) >= 4 and sum(1 for c in bullet_counts if c == 3) / len(bullet_counts) > 0.6:
        _f(out, "W5", None, "More than 60% of slides have exactly 3 bullets (rule of three)")
    sigs = [_signature(sl, deck["height"]) for sl in deck["slides"]]
    run = 1
    for k in range(1, len(sigs)):
        run = run + 1 if sigs[k] and sigs[k] == sigs[k - 1] else 1
        if run == 3:
            _f(out, "L2", deck["slides"][k]["idx"], "Same layout on 3 slides in a row")


def _inherits_effects(sl, flags):
    """Is there a shape that references a theme effectStyle (index idx, from 1) with a shadow or glow and does not cut it with an effectLst."""
    for el in sl["els"]:
        ref = el.find(f"{qn('p:style')}/{qn('a:effectRef')}")
        spPr = el.find(qn("p:spPr"))
        if ref is None or (spPr is not None and spPr.find(qn("a:effectLst")) is not None):
            continue
        idx = int(ref.get("idx", "0"))
        if 1 <= idx <= len(flags) and flags[idx - 1]:
            return True
    return False


def _check_title_rule(sl, out):
    title = next((s for s in sl["shapes"] if s["is_title"]), None)
    if not title:
        return
    bottom = title["y"] + title["h"]
    for s in sl["shapes"]:
        flat = s["kind"] in ("line", "rect") and s["h"] < (0.05 if s["kind"] == "line" else 0.1)
        if (flat and not s["is_title"] and s["w"] > 0.5 and abs(s["x"] - title["x"]) < 0.5
                and bottom - 0.05 <= s["y"] <= bottom + 0.6):
            _f(out, "L6", sl["idx"], "Accent rule under the title")
            return


def _check_side_stripe(sl, out):
    cards = [s for s in sl["shapes"] if s["kind"] in ("rect", "text") and s["w"] >= 1 and s["h"] >= 0.6]
    for st in (s for s in sl["shapes"] if s["kind"] == "rect"):
        vertical = st["w"] < 0.15 and st["h"] > 0.4
        horizontal = st["h"] < 0.15 and st["w"] > 0.4
        for card in cards:
            if card is st:
                continue
            if vertical and abs(st["x"] - card["x"]) < 0.06 and abs(st["y"] - card["y"]) < 0.1 and abs(st["h"] - card["h"]) < 0.15:
                _f(out, "L5", sl["idx"], "Colored stripe on one side of a card")
                return
            if horizontal and abs(st["y"] - card["y"]) < 0.06 and abs(st["x"] - card["x"]) < 0.1 and abs(st["w"] - card["w"]) < 0.15:
                _f(out, "L5", sl["idx"], "Colored stripe on top of a card")
                return


def _check_fake_chart(sl, out):
    if any(s["kind"] == "chart" for s in sl["shapes"]):
        return
    bars = [s for s in sl["shapes"] if s["kind"] == "rect" and s["w"] > 0.1 and s["h"] > 0.1]
    groupings = (
        (lambda b: (round(b["w"], 1), round(b["y"] + b["h"], 1)), lambda b: b["h"]),  # vertical bars: same width and baseline
        (lambda b: (round(b["h"], 1), round(b["x"], 1)), lambda b: b["w"]),           # horizontal bars: same height and left edge
    )
    for key_fn, size_fn in groupings:
        groups = {}
        for b in bars:
            groups.setdefault(key_fn(b), []).append(b)
        if any(len(g) >= 3 and len({round(size_fn(b), 2) for b in g}) >= 3 for g in groups.values()):
            _f(out, "D5", sl["idx"], "Fake bar chart drawn with shapes: use a native chart")
            return


def _check_colors(sl, out):
    bg = re.search(r"<p:bg>.*?</p:bg>", sl["xml"], re.S)
    bg_colors = {c.upper() for c in re.findall(r'srgbClr val="([0-9A-Fa-f]{6})"', bg.group(0))} if bg else set()
    colors = {c.upper() for c in re.findall(r'srgbClr val="([0-9A-Fa-f]{6})"', sl["xml"])} - bg_colors
    if len(colors) > 3:
        _f(out, "C10", sl["idx"], f"{len(colors)} colors on one slide (excluding background)")


def check_xml(deck, style_fonts, out):
    for sl in deck["slides"]:
        i, xml = sl["idx"], sl["xml"]
        if "<a:gradFill" in xml:
            _f(out, "C1", i, "Gradient fill")
        if re.search(r"<a:(outerShdw|glow|innerShdw)\b", xml):
            _f(out, "C6", i, "Shadow/glow effect")
        elif _inherits_effects(sl, deck["effect_flags"]):
            _f(out, "C6", i, "Shape inheriting a theme shadow (no effectLst)")
        _check_title_rule(sl, out)
        _check_side_stripe(sl, out)
        _check_fake_chart(sl, out)
        _check_colors(sl, out)
        hangul_runs = [r for r in sl["runs"] if HANGUL.search(r["text"])]
        if any(r["ea"] in ("", "+mn-ea", "+mj-ea") for r in hangul_runs) and not deck["ko_font_set"] and not style_fonts & OFFICE_KO:
            _f(out, "T12", i, "Hangul run has no East Asian (ea) font set: will be substituted with Malgun Gothic/Gulim")
        if any(not r["lang"].lower().startswith("ko") for r in hangul_runs):
            _f(out, "T11", i, "Hangul run lacks lang=ko-KR: lines will break mid-word")
        used = {f.lower() for r in sl["runs"] for f in (r["latin"], r["ea"]) if f}
        for font in sorted(used):
            if font in FONT_RULE and font not in style_fonts:
                _f(out, FONT_RULE[font], i, f"AI-default font '{font}'")


def check_theme(deck, style_fonts, out):
    for font in sorted({f.lower() for f in deck["theme_fonts"] if f}):
        if font in FONT_RULE and font not in style_fonts:
            _f(out, FONT_RULE[font], None, f"Theme default font is '{font}'")


def lint(path, style=None):
    deck = load(path)
    style_fonts = set()
    if style:
        font = json.loads(Path(style).read_text(encoding="utf-8")).get("font", {})
        style_fonts = {str(v).lower() for v in font.values() if isinstance(v, str)}
    out = []
    check_text(deck, out)
    check_titles(deck, out)
    check_structure(deck, out)
    check_xml(deck, style_fonts, out)
    check_theme(deck, style_fonts, out)
    out.sort(key=lambda f: (ORDER[f["severity"]], f["slide"] or 0))
    counts = Counter(f["severity"] for f in out)
    notes = {str(sl["idx"]): sl["notes"] for sl in deck["slides"] if sl["notes"].strip()}  # the reviewer checks H5, S9, and sources that appear only in the notes
    return {"file": str(path), "counts": {k: counts.get(k, 0) for k in ORDER}, "notes": notes, "findings": out}


def main(argv=None):
    ap = argparse.ArgumentParser(description="Automated AI-tell check")
    ap.add_argument("pptx")
    ap.add_argument("--style")
    ap.add_argument("--out")
    a = ap.parse_args(argv)
    out = Path(a.out) if a.out else Path(a.pptx).with_name("lint.json")
    try:
        res = lint(a.pptx, a.style)
    except Exception as e:  # corrupt file, etc.: leaving the previous result would be mistaken for a pass
        out.unlink(missing_ok=True)
        print(f"Lint failed: {type(e).__name__}: {e}", file=sys.stderr)
        sys.exit(2)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(res, ensure_ascii=False, indent=1), encoding="utf-8")
    for f in res["findings"]:
        where = f"slide {f['slide']}" if f["slide"] else "whole deck"
        print(f"[{f['severity']}] ({f['rule']}) {where}: {f['message']}")
    print(f"Total {res['counts']} -> {out}")
    sys.exit(1 if res["counts"]["blocker"] + res["counts"]["major"] else 0)


if __name__ == "__main__":
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8")
    main()
