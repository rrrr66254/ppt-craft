"""Extract style properties (colors, fonts, sizes, margins) from a reference pptx. Raw material for a style.json draft.
Usage: python extract_pptx_style.py <ref.pptx>
Look at layout and motifs by capturing with render.py (this script only gives numbers)."""
import json
import re
import sys
from collections import Counter

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
_COLOR = re.compile(r'<a:srgbClr val="([0-9A-Fa-f]{6})"')
_FONT = re.compile(r'<a:(?:latin|ea) typeface="([^"+][^"]*)"')
_SIZE = re.compile(r'<a:(?:rPr|defRPr|endParaRPr)[^>]*\ssz="(\d+)"')


def _theme(prs):
    root = etree.fromstring(prs.slide_master.part.part_related_by(RT.THEME).blob)
    colors = {}
    for node in root.find(".//a:clrScheme", NS):
        if etree.QName(node).localname == "extLst" or not len(node):
            continue
        child = node[0]
        kind = etree.QName(child).localname
        if kind == "srgbClr":
            colors[etree.QName(node).localname] = f"#{child.get('val').upper()}"
        elif kind == "sysClr":
            colors[etree.QName(node).localname] = f"#{child.get('lastClr', '000000').upper()}"
        elif kind == "prstClr":  # preset color name (e.g. black)
            colors[etree.QName(node).localname] = child.get("val")
    fonts = {}
    for kind in ("majorFont", "minorFont"):
        node = root.find(f".//a:{kind}", NS)
        hang = next((f.get("typeface") for f in node.findall("a:font", NS) if f.get("script") == "Hang"), "")
        fonts[kind] = {"latin": node.find("a:latin", NS).get("typeface"),
                       "ea": node.find("a:ea", NS).get("typeface"), "hang": hang}
    return colors, fonts


def extract(path):
    prs = Presentation(str(path))
    colors, fonts = _theme(prs)
    used_colors, used_fonts, sizes = Counter(), Counter(), Counter()
    lefts = []
    for slide in prs.slides:
        xml = etree.tostring(slide._element, encoding="unicode")
        used_colors.update(f"#{c.upper()}" for c in _COLOR.findall(xml))
        used_fonts.update(_FONT.findall(xml))
        sizes.update(int(s) / 100 for s in _SIZE.findall(xml))
        lefts += [sh.left for sh in slide.shapes if sh.has_text_frame and sh.left and sh.left > 0]
    return {
        "slide_size_in": [round(prs.slide_width.inches, 3), round(prs.slide_height.inches, 3)],
        "slides": len(prs.slides),
        "theme_colors": colors,
        "theme_fonts": fonts,
        "used_colors": used_colors.most_common(10),
        "used_fonts": used_fonts.most_common(8),
        "font_sizes_pt": sizes.most_common(8),
        "margin_guess_in": round(min(lefts).inches, 2) if lefts else None,
    }


if __name__ == "__main__":
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8")
    print(json.dumps(extract(sys.argv[1]), ensure_ascii=False, indent=1))
